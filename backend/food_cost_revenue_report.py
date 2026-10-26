"""Food cost vs revenue PDF — Food reports page is source of truth for headcounts."""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle
import io

from db import db
from food_revenue_report import build_food_revenue_report, _inr, _num, DAY_LABELS
from util import to_ist, now_utc

EXPENSE_XLSX = Path(
    "/root/.cursor/projects/root-One-10-Events-One-10-Durga-Pujo-2026/uploads/"
    "9b3aaee5-aa52-412d-a907-7f50a754d958/Durga_Puja_Food_Expense (1).xlsx"
)
# Fallback to earlier upload if final missing
if not EXPENSE_XLSX.exists():
    EXPENSE_XLSX = Path(
        "/root/.cursor/projects/root-One-10-Events-One-10-Durga-Pujo-2026/uploads/"
        "f7d82b39-8dff-49fe-982b-7c8d345c8fae/Durga_Puja_Food_Expense.xlsx"
    )

DAY_MAP = {
    ("Sashthi", "16 Oct (Fri)"): "sasthi",
    ("Saptami", "17 Oct (Sat)"): "saptami",
    ("Saptami", "18 Oct (Sun)"): "saptami_ashtami",
    ("Ashtami", "19 Oct (Mon)"): "ashtami",
    ("Nabami", "20 Oct (Tue)"): "nabami",
    ("Dashami", "21 Oct (Wed)"): "dashami",
}
PACKET_ROUND = 10

VERMILION = colors.HexColor("#D9381E")
BROWN = colors.HexColor("#1F1412")
IVORY = colors.HexColor("#FDFBF7")
GOLD = colors.HexColor("#B5952F")
MUTED = colors.HexColor("#6B5E55")
GREEN = colors.HexColor("#047857")
AMBER = colors.HexColor("#B45309")


def _round_up(n: int, step: int) -> int:
    if not step or n <= 0:
        return n
    return ((n + step - 1) // step) * step


def _style(header_bg=BROWN, font=7):
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), header_bg),
        ("TEXTCOLOR", (0, 0), (-1, 0), IVORY if header_bg != GOLD else BROWN),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F0EBE4")),
        ("FONTSIZE", (0, 0), (-1, -1), font),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D4CBC3")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [IVORY, colors.HexColor("#FAF7F3")]),
    ])


async def load_food_reports_truth() -> dict:
    """Same payload as /admin/food-reports (confirmed + vendor = UI default)."""
    # Import inside to avoid circular import with routes_food.
    from routes_food import admin_food_reports
    user = {"user_id": "system", "name": "report", "roles": ["system_admin"]}
    return await admin_food_reports(user=user)


def index_truth(data: dict) -> dict:
    out = {}
    for m in data.get("meals") or []:
        key = (m["day_code"], m["meal_code"], m.get("diet") or "veg")
        conf = m.get("confirmed") or {}
        vend = m.get("vendor") or {}
        out[key] = {
            "heads": int(conf.get("heads") or 0) + int(vend.get("heads") or 0),
            "confirmed": int(conf.get("heads") or 0),
            "vendor": int(vend.get("heads") or 0),
            "kids": int(conf.get("kids") or 0) + int(vend.get("kids") or 0),
            "free": int(conf.get("free") or 0),
            "paid_plates": int(conf.get("paid_plates") or 0),
            "unit_paise": int(m.get("unit_paise") or 0),
        }
    return out


def truth_get(idx, day, meal, diet):
    return idx.get((day, meal, diet), {
        "heads": 0, "confirmed": 0, "vendor": 0, "kids": 0, "free": 0, "paid_plates": 0, "unit_paise": 0,
    })


async def build_expense_reconcile():
    data = await load_food_reports_truth()
    idx = index_truth(data)
    wb = load_workbook(str(EXPENSE_XLSX), data_only=True)
    ws = wb["Food Expense"]
    lines = []
    breakfast_rates = {}

    for row in ws.iter_rows(min_row=5, max_row=33, values_only=True):
        day, date, meal, rate, heads, expense = row
        if not meal:
            continue
        day_code = DAY_MAP[(day, date)]
        rate = float(rate or 0)
        excel_heads = float(heads or 0)
        excel_expense = float(expense or 0)
        m = meal.lower()
        kids = vendor = paid = free = 0
        note = ""

        if m == "breakfast":
            breakfast_rates[day_code] = rate
            reports = 0
            for diet in ("veg", "non_veg"):
                g = truth_get(idx, day_code, "breakfast", diet)
                reports += g["heads"]
                kids += g["kids"]
                vendor += g["vendor"]
                paid += g["paid_plates"]
                free += g["free"]
            note = "Food reports seating (confirmed+vendor)"
        elif "extra" in m:
            g = truth_get(idx, day_code, "lunch", "veg")
            g2 = truth_get(idx, day_code, "lunch", "non_veg")
            reports = g["heads"] + g2["heads"]
            kids = g["kids"] + g2["kids"]
            vendor = g["vendor"] + g2["vendor"]
            paid = g["paid_plates"] + g2["paid_plates"]
            free = g["free"] + g2["free"]
            note = f"Ashtami lunch ALL heads (veg {g['heads']}, NV {g2['heads']}); {paid} paid extras"
        elif "lunch" in m and "non" in m:
            g = truth_get(idx, day_code, "lunch", "non_veg")
            reports, kids, vendor, paid, free = g["heads"], g["kids"], g["vendor"], g["paid_plates"], g["free"]
        elif "lunch" in m:
            g = truth_get(idx, day_code, "lunch", "veg")
            reports, kids, vendor, paid, free = g["heads"], g["kids"], g["vendor"], g["paid_plates"], g["free"]
        elif "dinner" in m and "non" in m:
            g = truth_get(idx, day_code, "dinner", "non_veg")
            reports, kids, vendor, paid, free = g["heads"], g["kids"], g["vendor"], g["paid_plates"], g["free"]
        elif "dinner" in m:
            g = truth_get(idx, day_code, "dinner", "veg")
            reports, kids, vendor, paid, free = g["heads"], g["kids"], g["vendor"], g["paid_plates"], g["free"]
        else:
            reports = 0

        # Comp+vendor cost approximation: free plates + vendor heads (vendor already in reports heads)
        # Cost all reports heads at sheet rate (comps included — paid to vendor)
        cost_total = rate * reports
        # paid portion ≈ paid_plates; rest of (reports - paid) is comps/vendor/kids in heads
        cost_paid = rate * paid
        cost_comp = max(0, cost_total - cost_paid)

        lines.append({
            "day": day, "date": date, "day_code": day_code, "meal": meal, "rate": rate,
            "excel_heads": excel_heads, "reports_heads": reports,
            "delta": reports - excel_heads,
            "kids": kids, "vendor": vendor, "paid_heads": paid, "free": free,
            "excel_expense": excel_expense, "reports_cost": cost_total,
            "cost_paid": cost_paid, "cost_comp": cost_comp,
            "note": note,
            "is_ashtami": "extra" in m,
            "is_breakfast": m == "breakfast",
        })

    packet_lines = []
    packet_cost = packet_heads = 0
    for day_code, rate in breakfast_rates.items():
        actual = sum(truth_get(idx, day_code, "breakfast_packet", d)["heads"] for d in ("veg", "non_veg"))
        shown = _round_up(actual, PACKET_ROUND)
        paid = sum(truth_get(idx, day_code, "breakfast_packet", d)["paid_plates"] for d in ("veg", "non_veg"))
        cost = rate * shown
        packet_cost += cost
        packet_heads += shown
        packet_lines.append({
            "day_code": day_code, "label": DAY_LABELS.get(day_code, day_code),
            "actual": actual, "shown": shown, "paid": paid, "rate": rate, "cost": cost,
        })

    excel_total = sum(l["excel_expense"] for l in lines)
    reports_cost = sum(l["reports_cost"] for l in lines)
    cost_of_paid = sum(l["cost_paid"] for l in lines)
    cost_of_comp = sum(l["cost_comp"] for l in lines)

    bhog_paid = 0
    bhog_lines = []
    async for ln in db.bank_statement_lines.find({"type": "debit"}, {"_id": 0}):
        desc = ln.get("description") or ""
        if "bhog" in desc.lower() or "arunbhog" in desc.lower().replace(" ", ""):
            amt = int(ln["amount_paise"])
            bhog_paid += amt
            bhog_lines.append({
                "date": ln.get("date"), "amount_paise": amt,
                "matched": bool(ln.get("matched")), "desc": desc[:70],
            })

    return {
        "lines": lines,
        "excel_total_paise": int(round(excel_total * 100)),
        "reports_cost_paise": int(round(reports_cost * 100)),
        "variance_paise": int(round((excel_total - reports_cost) * 100)),
        "cost_of_paid_paise": int(round(cost_of_paid * 100)),
        "cost_of_comp_paise": int(round(cost_of_comp * 100)),
        "packet_cost_paise": int(round(packet_cost * 100)),
        "packet_heads": packet_heads,
        "packet_lines": packet_lines,
        "full_cost_paise": int(round((reports_cost + packet_cost) * 100)),
        "bhog_paid_paise": bhog_paid,
        "bhog_lines": bhog_lines,
        "excel_file": EXPENSE_XLSX.name,
        "truth_note": "Headcounts = Food reports (confirmed + vendor). Awaiting payment excluded.",
    }


def make_pdf(revenue: dict, expense: dict) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=landscape(A4))
    W, H = landscape(A4)
    margin = 11 * mm
    gen = to_ist(now_utc()).strftime("%d %b %Y, %I:%M %p IST")
    rt = revenue["totals"]
    rev = rt["revenue_paise"]
    excel = expense["excel_total_paise"]
    cost = expense["reports_cost_paise"]
    var = expense["variance_paise"]
    cost_paid = expense["cost_of_paid_paise"]
    cost_comp = expense["cost_of_comp_paise"]
    pkt = expense["packet_cost_paise"]
    full = expense["full_cost_paise"]

    def band(sub=""):
        c.setFillColor(VERMILION)
        c.rect(0, H - 15 * mm, W, 15 * mm, fill=1, stroke=0)
        c.setFillColor(IVORY)
        c.setFont("Helvetica-Bold", 13)
        c.drawString(margin, H - 6.5 * mm, "One 10 Durga Pujo 2026")
        c.setFont("Helvetica", 8)
        c.drawString(margin, H - 11.5 * mm,
                     "Food cost vs revenue  ·  Food reports = source of truth  ·  comps costed to vendor")
        c.drawRightString(W - margin, H - 6.5 * mm, gen)
        if sub:
            c.drawRightString(W - margin, H - 11.5 * mm, sub)

    # Page 1
    band("Page 1 · P&L + variance")
    y = H - 22 * mm
    cards = [
        ("REVENUE (ex packets)", _inr(rev), f"{_num(rt['paid_plates'])} paid plates", GREEN),
        ("VENDOR COST (Food reports)", _inr(cost), "Confirmed+vendor × sheet rates", VERMILION),
        ("EXCEL SHEET TOTAL", _inr(excel), expense["excel_file"][:28], AMBER),
        ("VARIANCE (Excel − reports)", _inr(var), "Near-zero = kitchen ≈ portal", GREEN if abs(var) < 50000 else VERMILION),
    ]
    cw = (W - 2 * margin - 9 * mm) / 4
    for i, (lab, val, sub, col) in enumerate(cards):
        x = margin + i * (cw + 3 * mm)
        c.setFillColor(colors.white)
        c.setStrokeColor(colors.HexColor("#E8E0D8"))
        c.roundRect(x, y - 26 * mm, cw, 28 * mm, 3 * mm, fill=1, stroke=1)
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 6.5)
        c.drawString(x + 2.5 * mm, y - 3.5 * mm, lab)
        c.setFillColor(col)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x + 2.5 * mm, y - 11 * mm, val)
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 6)
        c.drawString(x + 2.5 * mm, y - 17 * mm, sub[:42])

    y = y - 34 * mm
    c.setFillColor(BROWN)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(margin, y, "Cost vs revenue")
    y -= 4 * mm
    pnl = [
        ["Item", "Amount", "Notes"],
        ["Food coupon revenue (ex breakfast packets)", _inr(rev), "Paid resident plates only"],
        ["Vendor cost — paid plates @ sheet rates", _inr(cost_paid), "Residents paid for these"],
        ["Vendor cost — complimentary + vendor meals", _inr(cost_comp), "Still payable to caterer"],
        ["Vendor cost — Food reports seating/L/D", _inr(cost), "Source of truth headcounts"],
        ["Expense Excel total", _inr(excel), expense["excel_file"]],
        ["Variance Excel − Food reports cost", _inr(var), "Was ₹54k when wrong portal extract used"],
        ["+ Breakfast packets (round-up ×10) @ bfast rates", _inr(pkt),
         f"{_num(expense['packet_heads'])} packet heads — not on Excel"],
        ["Full vendor cost incl. packets", _inr(full), "Reports L/D/B + packets"],
        ["Margin: revenue − reports cost", _inr(rev - cost), "Comps included in cost"],
        ["Margin: revenue − full (incl. packets)", _inr(rev - full), ""],
        ["Margin: revenue − Excel", _inr(rev - excel), ""],
        ["Arun Bhog on ICICI so far", _inr(expense["bhog_paid_paise"]), f"{len(expense['bhog_lines'])} IMPS"],
    ]
    t = Table(pnl, colWidths=[105 * mm, 32 * mm, 120 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BROWN),
        ("TEXTCOLOR", (0, 0), (-1, 0), IVORY),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("FONTNAME", (1, 1), (1, 1), "Helvetica-Bold"),
        ("TEXTCOLOR", (1, 1), (1, 1), GREEN),
        ("BACKGROUND", (0, 6), (-1, 6), colors.HexColor("#ECFDF5")),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D4CBC3")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [IVORY, colors.HexColor("#FAF7F3")]),
    ]))
    _, th = t.wrapOn(c, W, H)
    t.drawOn(c, margin, y - th)

    y = y - th - 7 * mm
    c.setFillColor(BROWN)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(margin, y, "What was wrong earlier / what’s true now")
    y -= 4.5 * mm
    bullets = [
        "Earlier recon built its own headcounts (and dropped kids / mismatched meals). Ashtami “93 NV lunch” was actually Ashtami dinner NV (93) — lunch is 439 veg / 0 NV.",
        "Source of truth is now the Food reports page: confirmed plates + vendor meals (awaiting payment OFF), same as the kitchen dashboard screenshot.",
        f"Excel vs Food reports on seating/L/D: variance only {_inr(var)} (Excel ₹{_num(excel//100)} vs reports-cost ₹{_num(cost//100)}).",
        "Complimentary plates are included in vendor cost (you pay the caterer). Paid extras on Ashtami lunch = 17 plates.",
        "Breakfast packets are on Food reports but not on the Excel — shown separately (~"
        + _inr(pkt) + " if billed at sheet breakfast rates with packet round-up).",
    ]
    c.setFont("Helvetica", 7.5)
    for b in bullets:
        words = b.split()
        line = ""
        for w in words:
            trial = (line + " " + w).strip()
            if c.stringWidth(trial, "Helvetica", 7.5) > W - 2 * margin - 4 * mm:
                c.drawString(margin + 2 * mm, y, "• " + line if line == words[0] or not line.startswith("•") else line)
                # simpler:
                break
        # draw wrapped
        line = ""
        first = True
        for w in words:
            trial = (line + " " + w).strip()
            if c.stringWidth(("• " if first else "  ") + trial, "Helvetica", 7.5) > W - 2 * margin:
                c.drawString(margin, y, ("• " if first else "  ") + line)
                y -= 3.8 * mm
                line = w
                first = False
            else:
                line = trial
        if line:
            c.drawString(margin, y, ("• " if first else "  ") + line)
            y -= 4.5 * mm

    c.setFillColor(MUTED)
    c.setFont("Helvetica", 6.5)
    c.drawString(margin, 7 * mm, expense["truth_note"])
    c.showPage()

    # Page 2 — line variance
    band("Page 2 · Excel vs Food reports")
    y = H - 20 * mm
    c.setFillColor(BROWN)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(margin, y, "Line variance (Δ = Food reports − Excel). Negative Δ = Excel counted more than portal.")
    y -= 3.5 * mm
    data = [["Day / meal", "Excel", "Reports", "Δ", "Kids", "Vend", "Rate", "Excel ₹", "Reports ₹", "Δ ₹"]]
    for l in expense["lines"]:
        data.append([
            f"{l['day'][:7]} {l['meal']}"[:28],
            _num(int(l["excel_heads"])),
            _num(int(l["reports_heads"])),
            f"{int(l['delta']):+d}",
            _num(int(l["kids"])),
            _num(int(l["vendor"])),
            _inr(int(l["rate"] * 100)),
            _inr(int(round(l["excel_expense"] * 100))),
            _inr(int(round(l["reports_cost"] * 100))),
            _inr(int(round((l["reports_cost"] - l["excel_expense"]) * 100))),
        ])
    data.append([
        "TOTAL",
        _num(int(sum(l["excel_heads"] for l in expense["lines"]))),
        _num(int(sum(l["reports_heads"] for l in expense["lines"]))),
        f"{int(sum(l['delta'] for l in expense['lines'])):+d}",
        "", "", "",
        _inr(excel), _inr(cost), _inr(cost - excel),
    ])
    tbl = Table(data, colWidths=[48 * mm, 14 * mm, 16 * mm, 12 * mm, 12 * mm, 12 * mm,
                                 14 * mm, 22 * mm, 24 * mm, 22 * mm])
    tbl.setStyle(_style(font=6))
    for i, l in enumerate(expense["lines"], start=1):
        if abs(l["delta"]) >= 3:
            tbl.setStyle(TableStyle([("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFF7ED"))]))
        if l["is_ashtami"]:
            tbl.setStyle(TableStyle([("BACKGROUND", (0, i), (-1, i), colors.HexColor("#ECFDF5"))]))
    _, th = tbl.wrapOn(c, W, H)
    tbl.drawOn(c, margin, max(12 * mm, y - th))
    c.setFillColor(MUTED)
    c.setFont("Helvetica", 6.5)
    c.drawString(margin, 7 * mm, "Amber = |Δ|≥3 heads. Green = Ashtami lunch (439 = 439). Kids are inside Food reports heads; Excel note says exclude kids under 7.")
    c.showPage()

    # Page 3
    band("Page 3 · Revenue, packets, bank")
    y = H - 20 * mm
    day_data = [["Day", "Paid", "Revenue"]]
    for d in revenue["by_day"]:
        day_data.append([d["label"], _num(d["paid"]), _inr(d["revenue_paise"])])
    day_data.append(["TOTAL", _num(rt["paid_plates"]), _inr(rev)])
    dt = Table(day_data, colWidths=[42 * mm, 22 * mm, 28 * mm])
    dt.setStyle(_style())
    dtw, dth = dt.wrapOn(c, W, H)
    c.setFillColor(BROWN)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(margin, y, "Revenue by day")
    dt.drawOn(c, margin, y - 4 * mm - dth)

    pkt_data = [["Day", "Actual", "Shown×10", "Rate", "Est. cost"]]
    for p in expense["packet_lines"]:
        pkt_data.append([
            p["label"], _num(p["actual"]), _num(p["shown"]),
            _inr(int(p["rate"] * 100)), _inr(int(round(p["cost"] * 100))),
        ])
    pkt_data.append(["TOTAL", "", _num(expense["packet_heads"]), "", _inr(pkt)])
    pt = Table(pkt_data, colWidths=[36 * mm, 18 * mm, 22 * mm, 18 * mm, 26 * mm])
    pt.setStyle(_style(header_bg=GOLD, font=6.5))
    ptw, pth = pt.wrapOn(c, W, H)
    x2 = margin + dtw + 10 * mm
    c.drawString(x2, y, "Breakfast packets (not on Excel)")
    pt.drawOn(c, x2, y - 4 * mm - pth)

    y2 = y - 4 * mm - max(dth, pth) - 8 * mm
    c.setFillColor(BROWN)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(margin, y2, "Material variances explained")
    y2 -= 4 * mm
    explain = [
        ("Sasthi breakfast Excel 210 vs Reports 240", "Excel under-counted seating by 30 (incl. 7 vendor)."),
        ("Saptami NV dinner Excel 119 vs Reports 111", "Excel 8 above portal — likely kitchen buffer / walk-in."),
        ("18 Oct NV lunch Excel 123 vs Reports 127", "Reports 4 higher (portal has more booked plates)."),
        ("Dashami NV lunch Excel 180 vs Reports 173", "Excel 7 above portal."),
        ("Ashtami lunch / dinner NV", "Lunch = 439 veg only. The 93 NV is dinner — not lunch."),
    ]
    c.setFont("Helvetica", 7.5)
    for title, body in explain:
        c.setFillColor(BROWN)
        c.setFont("Helvetica-Bold", 7.5)
        c.drawString(margin, y2, title)
        y2 -= 3.6 * mm
        c.setFont("Helvetica", 7.5)
        c.setFillColor(MUTED)
        c.drawString(margin + 3 * mm, y2, body)
        y2 -= 5 * mm

    y2 -= 2 * mm
    c.setFillColor(BROWN)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(margin, y2, "Arun Bhog on ICICI")
    y2 -= 3.5 * mm
    bhog = [["Date", "Amount", "Matched", "Narration"]]
    for b in expense["bhog_lines"]:
        bhog.append([str(b["date"] or ""), _inr(b["amount_paise"]),
                     "Yes" if b["matched"] else "No", b["desc"][:50]])
    bhog.append(["TOTAL", _inr(expense["bhog_paid_paise"]), "", ""])
    bt = Table(bhog, colWidths=[28 * mm, 28 * mm, 22 * mm, 100 * mm])
    bt.setStyle(_style(font=7))
    _, bth = bt.wrapOn(c, W, H)
    bt.drawOn(c, margin, max(10 * mm, y2 - bth))

    c.setFillColor(MUTED)
    c.setFont("Helvetica", 6.5)
    c.drawString(margin, 7 * mm, "https://one10events.in/uploads/one10_food_cost_revenue.pdf")
    c.showPage()
    c.save()
    return buf.getvalue()


async def generate_food_cost_revenue_pdf():
    revenue = await build_food_revenue_report()
    expense = await build_expense_reconcile()
    return make_pdf(revenue, expense), revenue, expense


def publish_pdf(pdf: bytes) -> list[Path]:
    outs = [
        Path("/tmp/one10_food_cost_revenue.pdf"),
        Path("/root/One 10 Events/One-10-Durga-Pujo-2026/frontend/public/reports/one10_food_cost_revenue.pdf"),
        Path("/root/One 10 Events/One-10-Durga-Pujo-2026/frontend/build/reports/one10_food_cost_revenue.pdf"),
        Path("/var/www/one10events/build/reports/one10_food_cost_revenue.pdf"),
        Path("/var/www/one10events/uploads/one10_food_cost_revenue.pdf"),
    ]
    for p in outs:
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(pdf)
    return outs
