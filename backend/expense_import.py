"""Excel template download + upload import for paid expenses."""
from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TEMPLATE_FILENAME = "one10_expense_template.xlsx"

TEMPLATE_HEADERS = [
    "Date",
    "Payee",
    "Purpose",
    "Category Code",
    "Amount (₹)",
    "Paid From",
    "Payment Mode",
    "UTR / reference",
    "Notes",
]

HEADER_ALIASES = {
    "date": "date",
    "expense_date": "date",
    "payee": "payee",
    "vendor": "payee",
    "paid_to": "payee",
    "purpose": "purpose",
    "description": "purpose",
    "category_code": "account_code",
    "account_code": "account_code",
    "category": "account_code",
    "amount": "amount_rupees",
    "amount_rs": "amount_rupees",
    "amount_rupees": "amount_rupees",
    "paid_from": "paid_from",
    "from": "paid_from",
    "payment_mode": "payment_mode",
    "mode": "payment_mode",
    "utr": "utr",
    "utr_reference": "utr",
    "reference": "utr",
    "notes": "notes",
    "note": "notes",
}


def _normalize_header(h: str) -> str:
    return re.sub(r"[^a-z0-9_]+", "_", (h or "").strip().lower()).strip("_")


def _map_header(h: str) -> str | None:
    key = _normalize_header(h)
    if key in HEADER_ALIASES:
        return HEADER_ALIASES[key]
    if key.startswith("amount"):
        return "amount_rupees"
    return None


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def build_template_xlsx(*, category_codes: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"

    header_fill = PatternFill("solid", fgColor="7A1F1F")
    header_font = Font(color="FFFFFF", bold=True)
    for col, title in enumerate(TEMPLATE_HEADERS, start=1):
        cell = ws.cell(1, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    samples = [
        ["2026-10-10", "Example Vendor", "Pujo decoration materials", "5100", 12500, "bank", "upi", "UTR111222333", ""],
        ["2026-10-11", "Example Caterer", "Committee tasting snacks", "5200", 2500, "cash", "cash", "", "Sample row — replace"],
    ]
    for r_idx, row in enumerate(samples, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(r_idx, c_idx, val)

    widths = [12, 22, 28, 14, 12, 12, 14, 18, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    codes = category_codes or ["5100", "5200", "5900"]
    code_list = ",".join(codes[:40])
    if code_list:
        dv_cat = DataValidation(type="list", formula1=f'"{code_list}"', allow_blank=False)
        ws.add_data_validation(dv_cat)
        dv_cat.add("D2:D5000")

    dv_from = DataValidation(type="list", formula1='"bank,cash"', allow_blank=False)
    ws.add_data_validation(dv_from)
    dv_from.add("F2:F5000")

    dv_mode = DataValidation(type="list", formula1='"upi,neft,cash,cheque"', allow_blank=False)
    ws.add_data_validation(dv_mode)
    dv_mode.add("G2:G5000")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(2, len(samples) + 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_rows(ws) -> list[list]:
    return [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]


def parse_import_rows(content: bytes, filename: str) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        grid = None
        for title in wb.sheetnames:
            if title.lower() in ("instructions", "readme", "help"):
                continue
            grid = _sheet_rows(wb[title])
            break
        if not grid:
            return []
    else:
        text = content.decode("utf-8-sig", errors="replace")
        grid = list(csv.reader(io.StringIO(text)))

    if not grid:
        return []

    header_idx = None
    mapped = []
    for i, raw in enumerate(grid[:30]):
        if not any(_cell_str(c) for c in raw):
            continue
        cols = [_map_header(_cell_str(c)) for c in raw]
        if "payee" in cols and "amount_rupees" in cols:
            header_idx = i
            mapped = cols
            break
    if header_idx is None:
        raise ValueError("Could not find header row. Use the downloaded template (Payee + Amount columns).")

    out = []
    for r_i, raw in enumerate(grid[header_idx + 1 :], start=header_idx + 2):
        if not any(_cell_str(c) for c in raw):
            continue
        row: dict[str, Any] = {"excel_row": r_i}
        for idx, key in enumerate(mapped):
            if not key:
                continue
            val = raw[idx] if idx < len(raw) else ""
            row[key] = _cell_str(val)
        # Skip untouched sample placeholders
        payee = (row.get("payee") or "").strip()
        if not payee:
            continue
        if payee.lower().startswith("example "):
            continue
        out.append(row)
    return out
