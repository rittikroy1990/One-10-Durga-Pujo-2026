"""Food subscriptions — public register + admin list (coupons/poll kept for compatibility)."""
import asyncio
import csv
import io
import re
from contextlib import asynccontextmanager
from copy import deepcopy
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, HTTPException, Body, Request, Response, UploadFile, File
from fastapi.responses import StreamingResponse
from pymongo.errors import DuplicateKeyError

from db import db, new_id, clean, NO_ID
from config import get_settings, get_active_cycle_id
from util import iso
from audit import audit, next_formatted
from auth import require, user_permissions, get_optional_user
from docs import food_coupon_pdf
from tokens import food_edit_token, food_voucher_token as _food_voucher_token, read_food_edit_token, read_food_voucher_token
from food_poll_catalog import build_catalog, DISHES, POLL_META, MEALS, DAYS

router = APIRouter(prefix="/api")

_DISH_IDS = {d["id"] for d in DISHES}
_DAY_CODES = {d["code"] for d in DAYS}
_MEAL_CODES = {m["code"] for m in MEALS}
_MAX_PER_MEAL = int(POLL_META.get("max_dishes_per_meal") or 8)
_ALLOWED_MEALS = {"breakfast", "lunch", "dinner"}


def _strip_evening_menu(food: dict) -> dict:
    """Public subscription menu: breakfast / lunch / dinner only."""
    out = deepcopy(food) if food else {}
    days = []
    for day in out.get("days") or []:
        meals = [m for m in (day.get("meals") or []) if (m.get("code") or "") in _ALLOWED_MEALS]
        if meals:
            d = dict(day)
            d["meals"] = meals
            days.append(d)
    out["days"] = days
    subtitle = out.get("subtitle") or ""
    for phrase in ("evening snacks", "Evening snacks", "evening snack", "Evening snack"):
        subtitle = subtitle.replace(phrase, "")
    subtitle = " ".join(subtitle.replace(" & ", " ").replace(" and ", " ").split())
    if subtitle:
        out["subtitle"] = subtitle
    return out


def _menu_lookup(food: dict) -> dict:
    """Map day_code|meal_code → meal meta from config (B/L/D only)."""
    food = _apply_default_prices(food)
    out = {}
    for day in food.get("days") or []:
        for meal in day.get("meals") or []:
            out[f"{day['code']}|{meal['code']}"] = {
                "day_code": day["code"],
                "day_label": day["label"],
                "meal_code": meal["code"],
                "meal_label": meal["label"],
                "amount_paise": meal.get("amount_paise"),
                "amount_label": meal.get("amount_label") or "TBC",
                "amount_status": meal.get("amount_status") or "TBC",
            }
    return out


def _normalize_selections(food: dict, selections: list, items_by_id: dict | None = None) -> list:
    """Accept cart lines with quantities.

    Supported shapes:
      - "day|meal"
      - {"day_code","meal_code", quantity|qty}
      - {"item_id", quantity|qty}  (menu catalog items)
    Quantities for the same key are merged. Max 50 per line.
    """
    items_by_id = items_by_id or {}
    lookup = _menu_lookup(food)
    cleaned = []
    index_by_key = {}
    for raw in selections or []:
        qty = 1
        item = None
        key = None
        if isinstance(raw, str) and "|" in raw:
            key = raw.strip()
            if key not in lookup:
                continue
            item = dict(lookup[key])
        elif isinstance(raw, dict) and (raw.get("item_id") or raw.get("id")):
            iid = (raw.get("item_id") or raw.get("id") or "").strip()
            src = items_by_id.get(iid)
            if not src or not src.get("active", True):
                continue
            try:
                qty = int(raw.get("quantity") if raw.get("quantity") is not None else raw.get("qty") or 1)
            except (TypeError, ValueError):
                qty = 1
            key = f"item|{iid}"
            item = {
                "item_id": iid,
                "name": src.get("name") or "",
                "description": src.get("description") or "",
                "image_url": _light_menu_image(src.get("image_url") or ""),
                "category": src.get("category") or "",
                "category_label": MEAL_CATEGORIES.get(src.get("category") or "") or src.get("category") or "",
                "diet": src.get("diet") or "",
                "diet_label": DIET_TYPES.get(src.get("diet") or "") or "",
                "menu_date": src.get("menu_date") or "",
                "item_day_code": src.get("day_code") or "",
                "day_code": "menu",
                "day_label": src.get("menu_date") or "Menu",
                "meal_code": src.get("category") or "item",
                "meal_label": src.get("name") or "Item",
                "amount_paise": int(src.get("amount_paise") or 0),
                "amount_label": src.get("amount_label") or "",
                "amount_status": "fixed",
            }
        elif isinstance(raw, dict):
            key = f"{raw.get('day_code')}|{raw.get('meal_code')}"
            try:
                qty = int(raw.get("quantity") if raw.get("quantity") is not None else raw.get("qty") or 1)
            except (TypeError, ValueError):
                qty = 1
            if key not in lookup:
                continue
            item = dict(lookup[key])
        else:
            continue
        if not item or qty < 1:
            continue
        qty = min(qty, 50)
        try:
            kids = max(0, min(20, int(raw.get("kid_qty") or 0))) if isinstance(raw, dict) else 0
        except (TypeError, ValueError):
            kids = 0
        if key in index_by_key:
            row = cleaned[index_by_key[key]]
            row["quantity"] = min(50, int(row.get("quantity") or 0) + qty)
            row["kid_qty"] = min(20, int(row.get("kid_qty") or 0) + kids)
            unit = row.get("amount_paise")
            row["line_total_paise"] = (int(unit) * row["quantity"]) if unit not in (None, "") else None
            continue
        item["quantity"] = qty
        item["kid_qty"] = kids
        unit = item.get("amount_paise")
        item["line_total_paise"] = (int(unit) * qty) if unit not in (None, "") else None
        index_by_key[key] = len(cleaned)
        cleaned.append(item)
    return cleaned


def _apply_complimentary(selections: list, used: dict, *, eligible: bool = True) -> dict:
    """Make the first N units of quota meals ₹0 (N per subscribed flat, minus already used)."""
    from pujo_menu_2026 import COMPLIMENTARY_QUOTAS, complimentary_key
    remaining = {
        k: (max(0, q - int((used or {}).get(k) or 0)) if eligible else 0)
        for k, q in COMPLIMENTARY_QUOTAS.items()
    }
    granted: dict = {}
    for sel in selections or []:
        cat = sel.get("category") or sel.get("meal_code") or ""
        day = sel.get("item_day_code") or sel.get("day_code") or ""
        key = complimentary_key(cat, day)
        qty = int(sel.get("quantity") or 0)
        free = min(remaining.get(key, 0), qty) if key else 0
        sel["complimentary_qty"] = free
        unit = sel.get("amount_paise")
        if unit not in (None, ""):
            sel["line_total_paise"] = int(unit) * (qty - free)
        if free:
            remaining[key] -= free
            granted[key] = granted.get(key, 0) + free
    return granted


# Unpaid orders with no payment proof give their free units back after this long.
COMPLIMENTARY_HOLD_HOURS = 12
_RELEASED_PAYMENT_STATUSES = ["expired", "cancelled"]
_flat_locks: dict[str, asyncio.Lock] = {}


def _flat_lock(flat_id: str) -> asyncio.Lock:
    lock = _flat_locks.get(flat_id)
    if lock is None:
        lock = _flat_locks[flat_id] = asyncio.Lock()
    return lock


@asynccontextmanager
async def _flat_guard(flat_id: str):
    """Serialise free-meal claims for one flat across requests and across server processes."""
    async with _flat_lock(flat_id):
        key, token = f"food_free:{flat_id}", new_id("lk")
        for _ in range(100):
            now = datetime.now(timezone.utc)
            try:
                if await db.food_flat_locks.find_one_and_update(
                    {"_id": key, "until": {"$lt": now}},
                    {"$set": {"until": now + timedelta(seconds=20), "token": token}},
                    upsert=True,
                ) is not None or await db.food_flat_locks.find_one({"_id": key, "token": token}):
                    break
            except DuplicateKeyError:
                pass
            await asyncio.sleep(0.1)
        else:
            raise HTTPException(status_code=503, detail="Busy saving another order for this flat — please try again.")
        try:
            yield
        finally:
            await db.food_flat_locks.delete_one({"_id": key, "token": token})


async def _expire_abandoned_food_orders(flat_id: str) -> int:
    """Expire this flat's unpaid orders that never got a payment proof, releasing their free units."""
    cutoff = (datetime.now(timezone.utc) - timedelta(hours=COMPLIMENTARY_HOLD_HOURS)).isoformat()
    expired = 0
    async for o in db.food_subscriptions.find(
        {"flat_id": flat_id, "payment_status": "pending", "created_at": {"$lt": cutoff}, "whatsapp_posted_at": None},
        {"_id": 0, "id": 1, "intent_id": 1},
    ):
        intent = await db.subscription_intents.find_one({"id": o.get("intent_id")}, {"_id": 0, "status": 1}) \
            if o.get("intent_id") else None
        if intent and intent.get("status") != "payment_pending":
            continue
        if o.get("intent_id") and (
            await db.upi_submissions.find_one({"intent_id": o["intent_id"], "status": {"$ne": "rejected"}}, {"_id": 1})
            or await db.receipts.find_one({"intent_id": o["intent_id"]}, {"_id": 1})
        ):
            continue
        if intent:
            await db.subscription_intents.update_one(
                {"id": o["intent_id"], "status": "payment_pending"},
                {"$set": {"status": "superseded", "superseded_at": iso(), "supersede_reason": "food_order_expired_unpaid"}},
            )
        await db.food_subscriptions.update_one(
            {"id": o["id"]},
            {"$set": {"payment_status": "expired", "status": "expired", "updated_at": iso()}},
        )
        expired += 1
    return expired


async def _complimentary_used(flat_id: str) -> dict:
    """Free units held by this flat on every order that is still live (paid, free, or awaiting payment)."""
    if not flat_id:
        return {}
    used: dict = {}
    async for o in db.food_subscriptions.find(
        {"flat_id": flat_id, "payment_status": {"$nin": _RELEASED_PAYMENT_STATUSES}},
        {"_id": 0, "complimentary": 1},
    ):
        for k, v in (o.get("complimentary") or {}).items():
            used[k] = used.get(k, 0) + int(v or 0)
    return used


def _mobile10(value) -> str:
    d = re.sub(r"\D", "", str(value or ""))
    return d[-10:] if len(d) >= 10 else ""


def _mobile_input(value) -> str:
    """A typed or autofilled mobile: 10 digits, optionally with a +91 / 0 prefix. Anything else is ''."""
    d = re.sub(r"\D", "", str(value or ""))
    if len(d) == 12 and d.startswith("91"):
        d = d[2:]
    elif len(d) == 11 and d.startswith("0"):
        d = d[1:]
    return d if len(d) == 10 else ""


def _cart_signature(selections: list) -> list:
    return sorted(
        (s.get("item_id") or f"{s.get('day_code')}|{s.get('meal_code')}", int(s.get("quantity") or 0), int(s.get("kid_qty") or 0))
        for s in selections or []
    )


async def _flat_subscribed(flat_id: str) -> bool:
    """Complimentary meals need a paid Puja subscription for the flat (tower + flat).

    A flat whose Puja payment was taken as a donation counts as subscribed once its donations reach the
    subscription amount.
    """
    if not flat_id:
        return False
    households = [h["id"] async for h in db.households.find(
        {"flat_id": flat_id, "is_deleted": {"$ne": True}}, {"_id": 0, "id": 1},
    )]
    if not households:
        return False
    receipts = await db.receipts.find(
        {"household_id": {"$in": households}, "kind": {"$in": ["subscription", "donation"]},
         "status": "issued", "cycle_id": await get_active_cycle_id()},
        {"_id": 0, "kind": 1, "total_amount": 1},
    ).to_list(50)
    if any(r.get("kind") == "subscription" for r in receipts):
        return True
    s = await get_settings()
    sub_amount = int((s.get("subscription") or {}).get("base_amount_paise") or 0)
    donated = sum(int(r.get("total_amount") or 0) for r in receipts if r.get("kind") == "donation")
    return bool(sub_amount and donated >= sub_amount)


_PLACEHOLDER_NAMES = {"check", "test", "example", "n/a", "na", "-", "."}
_PLACEHOLDER_MOBILES = {"9876543210", "0000000000", "1234567890"}


def _real_contact(name, mobile) -> dict:
    """{name, mobile} when both look like a real person's details (not seed/import placeholders)."""
    name = (name or "").strip()
    mobile = _mobile10(mobile)
    if not name or name.lower() in _PLACEHOLDER_NAMES or not mobile or mobile in _PLACEHOLDER_MOBILES:
        return {}
    return {"name": name, "mobile": mobile}


async def _flat_contact(flat_id: str) -> dict:
    """Name and mobile already on record for a flat: its Puja payer first, then its latest food order."""
    households = [h async for h in db.households.find(
        {"flat_id": flat_id, "is_deleted": {"$ne": True}},
        {"_id": 0, "id": 1, "primary_name": 1, "primary_mobile": 1, "source": 1},
    )]
    if households:
        async for r in db.receipts.find(
            {"household_id": {"$in": [h["id"] for h in households]}, "status": "issued",
             "kind": {"$in": ["subscription", "donation"]}},
            {"_id": 0, "payer_name": 1, "payer_mobile": 1},
        ).sort("issued_at", -1):
            if contact := _real_contact(r.get("payer_name"), r.get("payer_mobile")):
                return contact
        for h in sorted(households, key=lambda h: h.get("source") == "food_subscription"):
            if contact := _real_contact(h.get("primary_name"), h.get("primary_mobile")):
                return contact
    async for o in db.food_subscriptions.find({"flat_id": flat_id}, {"_id": 0, "name": 1, "mobile": 1}).sort("created_at", -1):
        if contact := _real_contact(o.get("name"), o.get("mobile")):
            return contact
    return {}


async def _complimentary_remaining(flat_id: str, *, eligible: bool = True, giving_back: dict | None = None) -> dict:
    from pujo_menu_2026 import COMPLIMENTARY_QUOTAS
    if not eligible:
        return {k: 0 for k in COMPLIMENTARY_QUOTAS}
    used = await _complimentary_used(flat_id)
    for k, v in (giving_back or {}).items():
        used[k] = max(0, int(used.get(k) or 0) - int(v or 0))
    return {k: max(0, q - int(used.get(k) or 0)) for k, q in COMPLIMENTARY_QUOTAS.items()}


def _apply_kid_plates(selections: list, raw_kids) -> tuple[list, int]:
    """Validate the kids' years of birth and cap each line's ₹0 kid plates.

    Each kid below the age limit may add one plate per paid line; quota meals (breakfast, Ashtami
    lunch) never take kid plates — kids there are booked like everyone else.
    """
    from pujo_menu_2026 import KIDS_FREE_BELOW_AGE, complimentary_key, kid_age
    kids = []
    for k in (raw_kids if isinstance(raw_kids, list) else [])[:10]:
        year = k.get("birth_year") if isinstance(k, dict) else k
        if year in (None, ""):
            continue
        age = kid_age(year)
        if age is None:
            raise HTTPException(status_code=400, detail="Enter a valid year of birth for each child.")
        kids.append({"birth_year": int(year), "age": age})
    allowed = sum(1 for k in kids if k["age"] < KIDS_FREE_BELOW_AGE)
    total = 0
    for sel in selections:
        quota_meal = complimentary_key(sel.get("category") or sel.get("meal_code") or "",
                                       sel.get("item_day_code") or sel.get("day_code") or "")
        sel["kid_qty"] = 0 if quota_meal else min(int(sel.get("kid_qty") or 0), allowed)
        total += sel["kid_qty"]
    return kids, total


def _selection_totals(selections: list) -> tuple:
    """Return (total_paise|None, item_count)."""
    total_paise = 0
    item_count = 0
    for sel in selections or []:
        qty = max(1, int(sel.get("quantity") or 1))
        item_count += qty
        unit = sel.get("amount_paise")
        line = sel.get("line_total_paise")
        if line not in (None, ""):
            total_paise += int(line)
        elif unit not in (None, ""):
            total_paise += int(unit) * qty
        else:
            return None, item_count
    return total_paise, item_count


DEFAULT_MEAL_PRICES_PAISE = {
    "breakfast": 7000,   # ₹70
    "lunch": 30000,      # ₹300
    "dinner": 30000,     # ₹300
}


def _apply_default_prices(food: dict) -> dict:
    """Ensure B/L/D meals have amount_paise + ₹ labels when missing."""
    out = _strip_evening_menu(food or {})
    for day in out.get("days") or []:
        for meal in day.get("meals") or []:
            code = meal.get("code") or ""
            if meal.get("amount_paise") in (None, "", 0) and code in DEFAULT_MEAL_PRICES_PAISE:
                meal["amount_paise"] = DEFAULT_MEAL_PRICES_PAISE[code]
            paise = meal.get("amount_paise")
            if paise not in (None, ""):
                try:
                    rupees = int(paise) // 100
                    meal["amount_label"] = f"₹{rupees}"
                    meal["amount_status"] = "fixed"
                except Exception:
                    pass
    return out


def _meal_price_map(food: dict) -> dict:
    food = _apply_default_prices(food)
    prices = {}
    for day in food.get("days") or []:
        for meal in day.get("meals") or []:
            code = meal.get("code")
            if code and code not in prices:
                prices[code] = {
                    "code": code,
                    "label": meal.get("label") or code.title(),
                    "amount_paise": meal.get("amount_paise"),
                    "amount_label": meal.get("amount_label") or "TBC",
                }
    return prices



def _food_page_enabled(food: dict) -> bool:
    """Public Food page is coming-soon unless admin explicitly enables it."""
    return bool((food or {}).get("page_enabled"))


# Once ordering is closed to the community, only committee members (logged in to the admin) can still order.
_FOOD_COMMITTEE_PERMS = ("households:write", "ops:manage", "receipts:manage", "expense:create")
_DEFAULT_CLOSED_MESSAGE = "Food orders for Durgotsav 2026 are now closed. Thank you for ordering with us."


def _community_closed(food: dict) -> bool:
    return bool((food or {}).get("community_closed"))


def _closed_message(food: dict) -> str:
    return ((food or {}).get("closed_message") or "").strip() or _DEFAULT_CLOSED_MESSAGE


async def _committee_orderer(request: Request | None) -> dict | None:
    user = await get_optional_user(request) if request is not None else None
    if user and any(p in user_permissions(user) for p in _FOOD_COMMITTEE_PERMS):
        return user
    return None

@router.get("/food/menu")
async def food_menu(request: Request):
    s = await get_settings()
    food = _apply_default_prices(s.get("food_subscription") or {})
    page_enabled = _food_page_enabled(food)
    closed = page_enabled and _community_closed(food)
    committee = await _committee_orderer(request) if closed else None
    if closed and not committee:
        return {"page_enabled": False, "closed": True, "closed_message": _closed_message(food), "items": [],
                "menu": {"days": []}, "payment_enabled": False, "meal_prices": [], "mode": "items"}
    catalog = []
    if page_enabled:
        catalog = await db.food_menu_items.find(
            {"active": True, "is_deleted": {"$ne": True}}, NO_ID
        ).sort([("menu_date", 1), ("category", 1), ("sort_order", 1), ("created_at", 1)]).to_list(200)
    from pujo_menu_2026 import MENU_META
    return {
        "menu": food if page_enabled else {"days": []},
        "items": [_item_public(i) for i in catalog],
        "page_enabled": page_enabled,
        "coming_soon_message": (food.get("coming_soon_message") or "").strip()
            or "Food subscriptions will open soon. Please check back.",
        "payment_enabled": bool(food.get("payment_enabled")) if page_enabled else False,
        "payment_note": food.get("payment_note") or "Pay via UPI QR after registering.",
        "meal_prices": list(_meal_price_map(food).values()) if page_enabled else [],
        "mode": "items" if catalog else "meals",
        "overall_menu_url": (food.get("overall_menu_url") or "").strip() if page_enabled else "",
        "overall_menu_filename": (food.get("overall_menu_filename") or "").strip() if page_enabled else "",
        "menu_meta": MENU_META if page_enabled else {},
        "kids_note": (food.get("kids_note") or MENU_META.get("kids_note") or "") if page_enabled else "",
        "timings": (food.get("timings") or MENU_META.get("timings") or {}) if page_enabled else {},
        "committee_mode": bool(committee),
        "committee_name": (committee.get("name") or "") if committee else "",
    }


@router.get("/food/complimentary")
async def food_complimentary_remaining(flat_id: str = "", editing: str = ""):
    """Free units (e.g. breakfasts) this flat can still claim, and whether its contact is already on record.

    `editing` (an order's edit token) counts that order's free units as available, since checkout replaces it.
    """
    flat_id = (flat_id or "").strip()
    if not flat_id or not await db.flats.find_one({"id": flat_id}, {"_id": 1}):
        raise HTTPException(status_code=404, detail="Unknown flat.")
    await _expire_abandoned_food_orders(flat_id)
    subscribed = await _flat_subscribed(flat_id)
    contact = await _flat_contact(flat_id)
    return {
        "flat_id": flat_id,
        "subscribed": subscribed,
        "known": bool(contact),
        "known_name": (contact.get("name") or "").split()[0] if contact else "",
        "known_mobile_masked": _mask_mobile(contact.get("mobile")) if contact else "",
        "remaining": await _complimentary_remaining(flat_id, eligible=subscribed, giving_back=await _editing_free(editing, flat_id)),
    }


async def _editing_free(token: str, flat_id: str) -> dict:
    sid = read_food_edit_token((token or "").strip()) if token else None
    doc = await db.food_subscriptions.find_one(
        {"id": sid, "flat_id": flat_id, "payment_status": {"$nin": _RELEASED_PAYMENT_STATUSES}}, NO_ID,
    ) if sid else None
    if not doc or await _resident_change_problem(doc):
        return {}
    return doc.get("complimentary") or {}


@router.post("/food/register")
async def food_register(body: dict = Body(...), request: Request = None):
    """Public food order — menu items and/or day meals; UPI QR when enabled."""
    s = await get_settings()
    food = _apply_default_prices(s.get("food_subscription") or {})
    if not food:
        raise HTTPException(status_code=503, detail="Food subscription is not configured yet.")
    if not _food_page_enabled(food):
        raise HTTPException(status_code=403, detail="Food subscriptions are coming soon.")
    committee = await _committee_orderer(request) if _community_closed(food) else None
    if _community_closed(food) and not committee:
        raise HTTPException(status_code=403, detail=_closed_message(food))

    catalog = await db.food_menu_items.find(
        {"active": True, "is_deleted": {"$ne": True}}, NO_ID
    ).to_list(500)
    items_by_id = {c["id"]: c for c in catalog}
    raw_selections = body.get("selections") or []
    stale = [str(r.get("item_id") or r.get("id")) for r in raw_selections
             if isinstance(r, dict) and (r.get("item_id") or r.get("id")) and (r.get("item_id") or r.get("id")) not in items_by_id]
    if stale:
        gone = [g["name"] for g in await db.food_menu_items.find({"id": {"$in": stale}}, {"_id": 0, "name": 1}).to_list(50)
                if g.get("name")]
        what = f"{', '.join(gone)} {'is' if len(gone) == 1 else 'are'}" if gone else "Some items in your cart are"
        raise HTTPException(status_code=409, detail=(
            f"{what} no longer available. Refresh the page, check your cart and place the order again."))
    selections = _normalize_selections(food, raw_selections, items_by_id=items_by_id)
    if not selections:
        raise HTTPException(status_code=400, detail="Add at least one item to your cart.")

    tower_id = (body.get("tower_id") or "").strip()
    flat_id = (body.get("flat_id") or "").strip()
    if not tower_id or not flat_id:
        raise HTTPException(status_code=400, detail="Select your tower and flat.")
    flat = await db.flats.find_one({"id": flat_id, "tower_id": tower_id}, {"_id": 0, "number": 1})
    tower = await db.towers.find_one({"id": tower_id}, {"_id": 0, "name": 1})
    if not flat or not tower:
        raise HTTPException(status_code=400, detail="Select a valid tower and flat.")
    tower_name = tower.get("name") or ""
    flat_number = flat.get("number") or ""

    name = (body.get("name") or "").strip() or (await _flat_contact(flat_id)).get("name") or ""
    mobile = _mobile_input(body.get("mobile"))
    if not name:
        raise HTTPException(status_code=400, detail="Name is required.")
    if not re.fullmatch(r"[6-9]\d{9}", mobile):
        raise HTTPException(status_code=400, detail="Enter a valid 10-digit mobile number.")
    kids, kid_heads = _apply_kid_plates(selections, body.get("kids"))
    replaces_id = None
    if (body.get("replaces_token") or "").strip():
        replaces_id = read_food_edit_token(body["replaces_token"].strip())
        if not replaces_id:
            raise HTTPException(status_code=400, detail="This edit link has expired — open the order again from My orders.")

    # One flat's free-unit count and order insert must not interleave with another request.
    async with _flat_guard(flat_id):
        await _expire_abandoned_food_orders(flat_id)
        old = await db.food_subscriptions.find_one({"id": replaces_id}, NO_ID) if replaces_id else None
        if replaces_id and (not old or old.get("flat_id") != flat_id):
            raise HTTPException(status_code=409, detail="The order you are editing belongs to another flat.")
        if old and old.get("payment_status") in _RELEASED_PAYMENT_STATUSES:
            if old.get("replaced_by"):
                raise HTTPException(status_code=409, detail=(
                    f"Order #{old['id'][-6:].upper()} was already updated (new order "
                    f"#{old['replaced_by'][-6:].upper()}). Open My orders to see it."))
            raise HTTPException(status_code=409, detail=(
                f"Order #{old['id'][-6:].upper()} was cancelled or expired, so it can't be edited. "
                f"Start a new order instead."))
        if not old and (repeat := await _recent_identical_order(flat_id, mobile, selections)):
            return await _register_response(repeat, message="This order was already placed — here it is again.")
        if old and (problem := await _resident_change_problem(old)):
            raise HTTPException(status_code=409, detail=problem)
        subscribed = await _flat_subscribed(flat_id)
        used_before = await _complimentary_used(flat_id)
        for k, v in ((old or {}).get("complimentary") or {}).items():
            used_before[k] = max(0, used_before.get(k, 0) - int(v or 0))
        granted = _apply_complimentary(selections, used_before, eligible=subscribed)
        total_paise, item_count = _selection_totals(selections)
        fully_complimentary = bool(granted) and total_paise == 0

        payment_enabled = bool(food.get("payment_enabled")) and total_paise is not None and total_paise > 0
        cycle_id = await get_active_cycle_id()
        sid = new_id("food")
        doc = {
            "id": sid,
            "cycle_id": cycle_id,
            "name": name,
            "mobile": mobile,
            "tower_id": tower_id,
            "tower_name": tower_name,
            "flat_id": flat_id,
            "flat_number": flat_number,
            "family_members": int(body.get("family_members") or 1),
            "notes": (body.get("notes") or "").strip()[:500],
            "selections": selections,
            "selection_count": item_count,
            "kids": kids,
            "kid_heads": kid_heads,
            "amount_status": "fixed" if total_paise is not None else "TBC",
            "total_amount_paise": total_paise,
            "complimentary": granted,
            "flat_subscribed": subscribed,
            "payment_status": "complimentary" if fully_complimentary else ("pending" if payment_enabled else "not_open"),
            "status": "confirmed" if fully_complimentary else "registered",
            "coupon_id": None,
            "coupon_no": None,
            "intent_id": None,
            "created_at": iso(),
            "updated_at": iso(),
            "source": "committee" if committee else "public",
            "order_mode": "items" if any(s.get("item_id") for s in selections) else "meals",
            "replaces_order_id": old["id"] if old else None,
        }
        if committee:
            doc["placed_by"] = committee.get("user_id") or ""
            doc["placed_by_name"] = committee.get("name") or committee.get("login_id") or ""

        intent_id = None
        if payment_enabled:
            # One household per (cycle, tower, flat) is enforced by a unique index.
            existing_hh = await db.households.find_one(
                {"cycle_id": cycle_id, "tower_id": tower_id, "flat_id": flat_id}, {"_id": 0, "id": 1},
            )
            if existing_hh:
                hid = existing_hh["id"]
                await db.households.update_one({"id": hid}, {"$set": {"is_food_subscriber": True}})
            else:
                hid = new_id("hh")
                await db.households.insert_one({
                    "id": hid, "cycle_id": cycle_id,
                    "tower_id": tower_id,
                    "tower_name": tower_name,
                    "flat_id": flat_id,
                    "flat_number": flat_number,
                    "occupancy_type": "other", "family_members": doc["family_members"],
                    "primary_name": name, "primary_mobile": mobile or "",
                    "is_food_subscriber": True, "created_at": iso(), "is_deleted": False,
                    "source": "food_subscription",
                })
            intent_id = new_id("intent")
            intent = {
                "id": intent_id, "cycle_id": cycle_id, "household_id": hid,
                "kind": "food_subscription",
                "food_subscription_id": sid,
                "base_amount": total_paise, "donation_amount": 0, "total_amount": total_paise,
                "components": [],
                "payer_name": name, "payer_mobile": mobile or "",
                "status": "payment_pending", "method": "upi_qr",
                "notes": doc["notes"], "created_at": iso(),
            }
            await db.subscription_intents.insert_one(dict(intent))
            doc["intent_id"] = intent_id

        await db.food_subscriptions.insert_one(doc)
        if old:
            await db.food_subscriptions.update_one({"id": old["id"]}, {"$set": {"replaced_by": sid}})
            await _cancel_food_order(old, actor="public", reason=f"Replaced by edited order #{sid[-6:].upper()}",
                                     request=request)
    await audit(
        "food.subscription.register",
        actor=committee or "public",
        entity_type="food_subscription",
        entity_id=sid,
        after={
            "name": name,
            "mobile": (mobile[-4:] if mobile else ""),
            "selection_count": item_count,
            "total_amount_paise": total_paise,
            "payment_enabled": payment_enabled,
        },
        reason="Committee food order (closed to the community)" if committee else "Public food order",
        request=request,
    )
    if fully_complimentary:
        message = "Your complimentary meals are confirmed — nothing to pay. No take-aways."
    elif payment_enabled:
        message = "Order placed. Pay via UPI QR and upload your payment screenshot."
    else:
        message = "Order placed. Payment is not open yet."
    return await _register_response(doc, message=message, used_before=used_before,
                                    replaced_ref=old["id"][-6:].upper() if old else None)


# A second identical checkout from the same flat and mobile this soon is a double tap or a retry, not a new order.
_REPEAT_ORDER_WINDOW_SECONDS = 120


async def _recent_identical_order(flat_id: str, mobile: str, selections: list) -> dict | None:
    since = (datetime.now(timezone.utc) - timedelta(seconds=_REPEAT_ORDER_WINDOW_SECONDS)).isoformat()
    want = _cart_signature(selections)
    async for o in db.food_subscriptions.find(
        {"flat_id": flat_id, "mobile": mobile, "created_at": {"$gte": since},
         "payment_status": {"$in": ["pending", "not_open", "complimentary"]}, "replaced_by": None},
        NO_ID,
    ).sort("created_at", -1):
        if _cart_signature(o.get("selections")) == want:
            return o
    return None


async def _register_response(doc: dict, *, message: str, used_before: dict | None = None,
                             replaced_ref: str | None = None) -> dict:
    from tokens import status_token
    payment_enabled = bool(doc.get("intent_id"))
    total_paise = doc.get("total_amount_paise")
    return {
        "ok": True,
        "id": doc["id"],
        "ref": doc["id"][-6:].upper(),
        "status": doc["status"],
        "payment_status": doc["payment_status"],
        "payment_enabled": payment_enabled,
        "total_amount_paise": total_paise,
        "selection_count": doc.get("selection_count"),
        "kid_heads": doc.get("kid_heads") or 0,
        "selections": doc.get("selections") or [],
        "complimentary": doc.get("complimentary") or {},
        "complimentary_used_before": used_before or {},
        "flat_subscribed": doc.get("flat_subscribed"),
        "intent_id": doc.get("intent_id"),
        "status_token": status_token(doc["intent_id"]) if payment_enabled else None,
        "voucher_token": _food_voucher_token(doc["id"]) if doc.get("payment_status") == "complimentary" else None,
        "whatsapp_text": _whatsapp_order_text(doc, total_paise) if payment_enabled else None,
        "replaced_ref": replaced_ref,
        "message": message,
    }


@router.get("/admin/food-subscriptions/prices")
async def admin_food_prices(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    s = await get_settings()
    food = _apply_default_prices(s.get("food_subscription") or {})
    return {
        "page_enabled": _food_page_enabled(food),
        "community_closed": _community_closed(food),
        "closed_message": _closed_message(food),
        "coming_soon_message": (food.get("coming_soon_message") or "").strip(),
        "payment_enabled": bool(food.get("payment_enabled")),
        "payment_note": food.get("payment_note") or "",
        "prices": list(_meal_price_map(food).values()),
    }


MEAL_CATEGORIES = {
    "breakfast": "Breakfast",
    "breakfast_packet": "Breakfast packet",
    "lunch": "Lunch",
    "dinner": "Dinner",
}
_MEAL_SORT_TENS = {"breakfast": 10, "breakfast_packet": 20, "lunch": 30, "dinner": 40}
DIET_TYPES = {
    "veg": "Veg",
    "non_veg": "Non-veg",
}


def _normalize_meal_category(raw: str) -> str:
    key = (raw or "").strip().lower().replace("-", "_").replace(" ", "_")
    aliases = {
        "bf": "breakfast",
        "brkfast": "breakfast",
        "packet": "breakfast_packet",
        "breakfast_pack": "breakfast_packet",
        "l": "lunch",
        "d": "dinner",
        "din": "dinner",
    }
    key = aliases.get(key, key)
    return key if key in MEAL_CATEGORIES else ""


def _normalize_diet(raw: str) -> str:
    key = (raw or "").strip().lower().replace("-", "_").replace(" ", "_")
    aliases = {
        "vegetarian": "veg",
        "veggie": "veg",
        "nonveg": "non_veg",
        "nonvegetarian": "non_veg",
        "nv": "non_veg",
    }
    key = aliases.get(key, key)
    return key if key in DIET_TYPES else ""


def _normalize_menu_date(raw: str) -> str:
    """Accept YYYY-MM-DD (HTML date) or DD/MM/YYYY; return YYYY-MM-DD or ''."""
    s = (raw or "").strip()
    if not s:
        return ""
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return s
    m = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1 <= mo <= 12 and 1 <= d <= 31:
            return f"{y:04d}-{mo:02d}-{d:02d}"
    return ""


def _default_menu_image(category: str, diet: str, day_code: str = "") -> str:
    """Prefer day×meal×diet plate photos; fall back to generic meal samples."""
    cat = (category or "").strip().lower()
    if cat not in MEAL_CATEGORIES:
        return ""
    day = (day_code or "").strip().lower()
    diet_norm = (diet or "").strip().lower()
    diet_key = "nonveg" if diet_norm == "non_veg" else ("veg" if diet_norm == "veg" else "")
    if day and diet_key:
        return f"/images/food-menu/day/{day}-{cat}-{diet_key}.jpg"
    if not diet_key:
        return ""
    return f"/images/food-menu/menu-{diet_key}-{cat}.jpg"


def _light_menu_image(url: str) -> str:
    # The bundled plate photos also ship as ~80 KB JPEGs; the ~500 KB PNGs fail to load on slow phones.
    if url.startswith("/images/food-menu/") and url.endswith(".png"):
        return url[:-4] + ".jpg"
    return url


def _item_public(doc: dict) -> dict:
    category = _normalize_meal_category(doc.get("category") or "") or (doc.get("category") or "")
    diet = _normalize_diet(doc.get("diet") or doc.get("veg_type") or "")
    menu_date = _normalize_menu_date(doc.get("menu_date") or doc.get("date") or "") or (doc.get("menu_date") or "")
    day_code = (doc.get("day_code") or "").strip()
    image_url = _light_menu_image((doc.get("image_url") or "").strip() or _default_menu_image(category, diet, day_code))
    return clean({
        "id": doc.get("id"),
        "name": doc.get("name") or "",
        "description": doc.get("description") or "",
        "image_url": image_url,
        "amount_paise": int(doc.get("amount_paise") or 0),
        "amount_label": doc.get("amount_label") or f"₹{int(doc.get('amount_paise') or 0) // 100}",
        "category": category,
        "category_label": MEAL_CATEGORIES.get(category) or category or "",
        "diet": diet,
        "diet_label": DIET_TYPES.get(diet) or "",
        "menu_date": menu_date,
        "day_code": day_code,
        "day_label": doc.get("day_label") or "",
        "matrix_key": doc.get("matrix_key") or "",
        "active": bool(doc.get("active", True)),
        "sort_order": int(doc.get("sort_order") or 0),
        "is_complimentary": bool(doc.get("is_complimentary")),
        "complimentary_note": (doc.get("complimentary_note") or "").strip(),
        "price_note": (doc.get("price_note") or "").strip(),
        "badge": (doc.get("badge") or "").strip(),
        "for_all_diets": bool(doc.get("for_all_diets")),
        "created_at": doc.get("created_at"),
        "updated_at": doc.get("updated_at"),
    })


@router.get("/admin/food-menu-items")
async def admin_list_food_items(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    items = await db.food_menu_items.find(
        {"is_deleted": {"$ne": True}}, NO_ID
    ).sort([("sort_order", 1), ("created_at", -1)]).to_list(500)
    return {"items": [_item_public(i) for i in items]}


@router.post("/admin/food-menu-items")
async def admin_create_food_item(request: Request, user: dict = Depends(require("households:write", "ops:manage", "receipts:manage", "settings:manage"))):
    """Create a sellable menu item: date, meal category, veg/non-veg, image, price."""
    form = await request.form()
    name = (form.get("name") or "").strip()
    description = (form.get("description") or form.get("menu") or "").strip()[:800]
    category = _normalize_meal_category(form.get("category") or "")
    diet = _normalize_diet(form.get("diet") or form.get("veg_type") or "")
    menu_date = _normalize_menu_date(form.get("menu_date") or form.get("date") or "")
    price_raw = form.get("price_rupees") or form.get("price") or "0"
    try:
        price_rupees = float(str(price_raw).replace(",", "").strip() or 0)
    except ValueError:
        raise HTTPException(status_code=400, detail="Enter a valid price in rupees.")
    if not name or len(name) < 2:
        raise HTTPException(status_code=400, detail="Item name / menu title is required.")
    if not category:
        raise HTTPException(status_code=400, detail="Category must be Breakfast, Lunch, or Dinner.")
    if not diet:
        raise HTTPException(status_code=400, detail="Select Veg or Non-veg.")
    if not menu_date:
        raise HTTPException(status_code=400, detail="Menu date is required (YYYY-MM-DD).")
    if price_rupees < 0:
        raise HTTPException(status_code=400, detail="Price cannot be negative.")
    amount_paise = int(round(price_rupees * 100))
    active = str(form.get("active") or "true").lower() in ("1", "true", "yes", "on")
    try:
        sort_order = int(form.get("sort_order") or 0)
    except ValueError:
        sort_order = 0

    image_url = ""
    file = form.get("image")
    if file and getattr(file, "filename", None):
        data = await file.read()
        if data:
            if len(data) > 8 * 1024 * 1024:
                raise HTTPException(status_code=400, detail="Image too large (max 8 MB).")
            from storage import save_document
            ctype = getattr(file, "content_type", None) or "image/jpeg"
            doc = await save_document(
                data=data,
                filename=file.filename or "food-item.jpg",
                doc_type="food_menu",
                linked_type="food_menu_item",
                linked_id="pending",
                uploaded_by=user.get("user_id") or "admin",
                content_type=ctype,
            )
            image_url = doc.get("storage_path") or ""

    iid = new_id("fitem")
    if image_url:
        await db.documents.update_one({"storage_path": image_url}, {"$set": {"linked_id": iid}})
    rec = {
        "id": iid,
        "name": name,
        "description": description,
        "category": category,
        "diet": diet,
        "menu_date": menu_date,
        "image_url": image_url,
        "amount_paise": amount_paise,
        "amount_label": f"₹{int(amount_paise) // 100}",
        "active": active,
        "sort_order": sort_order,
        "is_deleted": False,
        "created_at": iso(),
        "updated_at": iso(),
        "created_by": user.get("user_id") or "",
    }
    await db.food_menu_items.insert_one(dict(rec))
    await audit("food.menu_item.create", entity_type="food_menu_item", entity_id=iid,
                after={"name": name, "amount_paise": amount_paise, "category": category, "diet": diet, "menu_date": menu_date},
                request=request)
    return {"ok": True, "item": _item_public(rec)}


@router.put("/admin/food-menu-items/{item_id}")
async def admin_update_food_item(item_id: str, request: Request, user: dict = Depends(require("households:write", "ops:manage", "receipts:manage", "settings:manage"))):
    existing = await db.food_menu_items.find_one({"id": item_id, "is_deleted": {"$ne": True}})
    if not existing:
        raise HTTPException(status_code=404, detail="Menu item not found.")
    form = await request.form()
    patch = {"updated_at": iso()}
    if "name" in form:
        name = (form.get("name") or "").strip()
        if len(name) < 2:
            raise HTTPException(status_code=400, detail="Item name is required.")
        patch["name"] = name
    if "description" in form or "menu" in form:
        patch["description"] = (form.get("description") or form.get("menu") or "").strip()[:800]
    if "category" in form:
        category = _normalize_meal_category(form.get("category") or "")
        if not category:
            raise HTTPException(status_code=400, detail="Category must be Breakfast, Lunch, or Dinner.")
        patch["category"] = category
    if "diet" in form or "veg_type" in form:
        diet = _normalize_diet(form.get("diet") or form.get("veg_type") or "")
        if not diet:
            raise HTTPException(status_code=400, detail="Select Veg or Non-veg.")
        patch["diet"] = diet
    if "menu_date" in form or "date" in form:
        menu_date = _normalize_menu_date(form.get("menu_date") or form.get("date") or "")
        if not menu_date:
            raise HTTPException(status_code=400, detail="Menu date is required (YYYY-MM-DD).")
        patch["menu_date"] = menu_date
    if "price_rupees" in form or "price" in form:
        try:
            price_rupees = float(str(form.get("price_rupees") or form.get("price") or 0).replace(",", ""))
        except ValueError:
            raise HTTPException(status_code=400, detail="Enter a valid price in rupees.")
        if price_rupees < 0:
            raise HTTPException(status_code=400, detail="Price cannot be negative.")
        amount_paise = int(round(price_rupees * 100))
        patch["amount_paise"] = amount_paise
        patch["amount_label"] = f"₹{amount_paise // 100}"
    if "active" in form:
        patch["active"] = str(form.get("active") or "").lower() in ("1", "true", "yes", "on")
    if "sort_order" in form:
        try:
            patch["sort_order"] = int(form.get("sort_order") or 0)
        except ValueError:
            pass
    file = form.get("image")
    if file and getattr(file, "filename", None):
        data = await file.read()
        if data:
            from storage import save_document
            ctype = getattr(file, "content_type", None) or "image/jpeg"
            doc = await save_document(
                data=data,
                filename=file.filename or "food-item.jpg",
                doc_type="food_menu",
                linked_type="food_menu_item",
                linked_id=item_id,
                uploaded_by=user.get("user_id") or "admin",
                content_type=ctype,
            )
            patch["image_url"] = doc.get("storage_path") or ""
    await db.food_menu_items.update_one({"id": item_id}, {"$set": patch})
    updated = await db.food_menu_items.find_one({"id": item_id}, NO_ID)
    await audit("food.menu_item.update", entity_type="food_menu_item", entity_id=item_id,
                after=patch, request=request)
    return {"ok": True, "item": _item_public(updated)}


@router.delete("/admin/food-menu-items/{item_id}")
async def admin_delete_food_item(item_id: str, request: Request = None, user: dict = Depends(require("households:write", "ops:manage", "receipts:manage", "settings:manage"))):
    existing = await db.food_menu_items.find_one({"id": item_id, "is_deleted": {"$ne": True}})
    if not existing:
        raise HTTPException(status_code=404, detail="Menu item not found.")
    await db.food_menu_items.update_one({"id": item_id}, {"$set": {
        "is_deleted": True, "active": False, "updated_at": iso(),
    }})
    await audit("food.menu_item.delete", entity_type="food_menu_item", entity_id=item_id, request=request)
    return {"ok": True}


@router.get("/admin/food-menu-matrix")
async def admin_food_menu_matrix(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    """Day × Meal × Veg/Non-veg matrix for admin pricing and images."""
    from food_menu_matrix import matrix_payload
    items = await db.food_menu_items.find({"is_deleted": {"$ne": True}}, NO_ID).to_list(500)
    payload = matrix_payload(items)
    s = await get_settings()
    food = s.get("food_subscription") or {}
    payload["page_enabled"] = _food_page_enabled(food)
    payload["payment_enabled"] = bool(food.get("payment_enabled"))
    payload["overall_menu_url"] = (food.get("overall_menu_url") or "").strip()
    payload["overall_menu_filename"] = (food.get("overall_menu_filename") or "").strip()
    payload["overall_menu_uploaded_at"] = food.get("overall_menu_uploaded_at") or ""
    return payload


@router.post("/admin/food-menu-matrix/seed-official")
async def admin_seed_official_pujo_menu(
    request: Request,
    user: dict = Depends(require("households:write", "ops:manage", "receipts:manage", "settings:manage")),
):
    """Replace/upsert every poster cell (incl. veg breakfast packets) + day plate images."""
    from pujo_menu_2026 import official_cells_for_seed, MENU_META
    cells = official_cells_for_seed()
    upserted = 0
    for cell in cells:
        key = cell["matrix_key"]
        existing = await db.food_menu_items.find_one({"matrix_key": key, "is_deleted": {"$ne": True}})
        patch = {
            **{k: cell[k] for k in (
                "matrix_key", "day_code", "day_label", "menu_date", "category", "diet",
                "name", "description", "amount_paise", "amount_label", "image_url",
                "active", "is_complimentary", "complimentary_note", "price_note", "badge", "sort_order",
                "for_all_diets",
            )},
            "updated_at": iso(),
            "is_deleted": False,
            "source": "official_poster_2026",
        }
        if existing:
            await db.food_menu_items.update_one({"id": existing["id"]}, {"$set": patch})
        else:
            await db.food_menu_items.insert_one({
                "id": new_id("fitem"),
                **patch,
                "created_at": iso(),
                "created_by": user.get("user_id") or "",
            })
        upserted += 1

    s = await get_settings()
    food = dict(s.get("food_subscription") or {})
    food["page_enabled"] = True
    food["payment_enabled"] = True if food.get("payment_enabled") is None else bool(food.get("payment_enabled"))
    food["kids_note"] = MENU_META.get("kids_note") or food.get("kids_note") or ""
    food["timings"] = MENU_META.get("timings") or food.get("timings") or {}
    food["overall_menu_url"] = food.get("overall_menu_url") or "/images/food-menu/pujo-menu-poster-2026-v2.jpg"
    food["overall_menu_filename"] = food.get("overall_menu_filename") or "Durga Puja Menu 2026 poster.jpg"
    food["payment_note"] = food.get("payment_note") or (
        "Pay the cart total via UPI QR after checkout, then upload your payment screenshot."
    )
    await db.application_settings.update_one({"id": "app_settings"}, {"$set": {"food_subscription": food}}, upsert=True)
    await audit(
        "food.menu_matrix.seed_official",
        entity_type="food_menu",
        entity_id="matrix",
        after={"cells": upserted},
        request=request,
    )
    return {"ok": True, "upserted": upserted, "message": f"Seeded {upserted} official menu cells from the Pujo poster."}


@router.put("/admin/food-menu-matrix/cell")
async def admin_upsert_matrix_cell(request: Request, user: dict = Depends(require("households:write", "ops:manage", "receipts:manage", "settings:manage"))):
    """Create/update one matrix cell (multipart): day_code, meal, diet, price, image, …"""
    from food_menu_matrix import (
        DAYS, MEALS, MATRIX_DIETS, matrix_key, default_cell_name, DEFAULT_CELL_PRICE_RUPEES,
        meal_allows_diet,
    )
    form = await request.form()
    day_code = (form.get("day_code") or "").strip().lower()
    meal = _normalize_meal_category(form.get("meal") or form.get("category") or "")
    diet = _normalize_diet(form.get("diet") or "")
    day = next((d for d in DAYS if d["code"] == day_code), None)
    if not day:
        raise HTTPException(status_code=400, detail="Unknown Pujo day code.")
    if not meal:
        raise HTTPException(status_code=400, detail="Meal must be breakfast, breakfast packet, lunch, or dinner.")
    if not diet:
        raise HTTPException(status_code=400, detail="Diet must be veg or non_veg.")
    if not meal_allows_diet(meal, diet):
        raise HTTPException(status_code=400, detail="Breakfast packet is pure veg only.")

    key = matrix_key(day_code, meal, diet)
    meal_label = next(m["label"] for m in MEALS if m["code"] == meal)
    diet_label = next(d["label"] for d in MATRIX_DIETS if d["code"] == diet)
    day_label = day.get("short_label") or day["label"]

    name = (form.get("name") or "").strip() or default_cell_name(day_label, meal_label, diet_label)
    description = (form.get("description") or form.get("menu") or "").strip()[:800]
    price_raw = form.get("price_rupees") or form.get("price")
    if price_raw is None or str(price_raw).strip() == "":
        price_rupees = float(DEFAULT_CELL_PRICE_RUPEES.get(meal, 300))
    else:
        try:
            price_rupees = float(str(price_raw).replace(",", "").strip())
        except ValueError:
            raise HTTPException(status_code=400, detail="Enter a valid price in rupees.")
    if price_rupees < 0:
        raise HTTPException(status_code=400, detail="Price cannot be negative.")
    amount_paise = int(round(price_rupees * 100))
    active = str(form.get("active") if form.get("active") is not None else "true").lower() in (
        "1", "true", "yes", "on", "y"
    )

    existing = await db.food_menu_items.find_one({"matrix_key": key, "is_deleted": {"$ne": True}})
    image_url = (existing or {}).get("image_url") or ""
    file = form.get("image")
    if file and getattr(file, "filename", None):
        data = await file.read()
        if data:
            if len(data) > 8 * 1024 * 1024:
                raise HTTPException(status_code=400, detail="Image too large (max 8 MB).")
            from storage import save_document
            ctype = getattr(file, "content_type", None) or "image/jpeg"
            doc = await save_document(
                data=data,
                filename=file.filename or "food-matrix.jpg",
                doc_type="food_menu",
                linked_type="food_menu_item",
                linked_id=(existing or {}).get("id") or "pending",
                uploaded_by=user.get("user_id") or "admin",
                content_type=ctype,
            )
            image_url = doc.get("storage_path") or image_url
    elif form.get("clear_image") in ("1", "true", "yes"):
        image_url = ""
    elif form.get("image_url"):
        image_url = str(form.get("image_url")).strip()[:500]

    patch = {
        "matrix_key": key,
        "day_code": day_code,
        "day_label": day_label,
        "menu_date": day["date"],
        "category": meal,
        "diet": diet,
        "name": name[:120],
        "description": description,
        "amount_paise": amount_paise,
        "amount_label": f"₹{amount_paise // 100}",
        "image_url": image_url,
        "active": active,
        "updated_at": iso(),
        "is_deleted": False,
        "is_complimentary": str(form.get("is_complimentary") or "").lower() in ("1", "true", "yes", "on", "y"),
        "complimentary_note": str(form.get("complimentary_note") or "").strip()[:400],
        "price_note": str(form.get("price_note") or "").strip()[:300],
        "badge": str(form.get("badge") or "").strip()[:120],
    }
    if "for_all_diets" in form:
        patch["for_all_diets"] = (
            diet == "veg" and str(form.get("for_all_diets")).lower() in ("1", "true", "yes", "on", "y")
        )

    if existing:
        await db.food_menu_items.update_one({"id": existing["id"]}, {"$set": patch})
        item_id = existing["id"]
        if image_url:
            await db.documents.update_one({"storage_path": image_url}, {"$set": {"linked_id": item_id}})
    else:
        item_id = new_id("fitem")
        rec = {
            "id": item_id,
            **patch,
            "sort_order": int(day.get("order") or 0) * 100 + _MEAL_SORT_TENS.get(meal, 40) + (1 if diet == "veg" else 2),
            "created_at": iso(),
            "created_by": user.get("user_id") or "",
            "source": "matrix",
        }
        await db.food_menu_items.insert_one(dict(rec))
        if image_url:
            await db.documents.update_one({"storage_path": image_url}, {"$set": {"linked_id": item_id}})

    updated = await db.food_menu_items.find_one({"id": item_id}, NO_ID)
    await audit("food.menu_matrix.cell_upsert", entity_type="food_menu_item", entity_id=item_id,
                after={"matrix_key": key, "amount_paise": amount_paise, "active": active}, request=request)
    return {"ok": True, "item": _item_public(updated)}


@router.get("/admin/food-menu-matrix/template")
async def admin_food_matrix_template(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    from food_menu_matrix import build_matrix_template_xlsx
    items = await db.food_menu_items.find({"is_deleted": {"$ne": True}}, NO_ID).to_list(500)
    data = build_matrix_template_xlsx(items)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="one10_food_menu_matrix.xlsx"'},
    )


@router.post("/admin/food-menu-matrix/import")
async def admin_food_matrix_import(
    file: UploadFile = File(...),
    request: Request = None,
    user: dict = Depends(require("households:write", "ops:manage", "receipts:manage", "settings:manage")),
):
    """Bulk upsert matrix cells from Excel/CSV (prices, names, active; optional image URL)."""
    from food_menu_matrix import parse_matrix_upload, MEALS, MATRIX_DIETS

    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="Empty file.")
    try:
        rows = parse_matrix_upload(raw, file.filename or "")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read Excel/CSV. Use the downloaded template.")

    created = updated = skipped = 0
    errors = []
    for row in rows:
        if row.get("error"):
            errors.append({"row": row.get("excel_row"), "error": row["error"]})
            skipped += 1
            continue
        key = row["matrix_key"]
        existing = await db.food_menu_items.find_one({"matrix_key": key, "is_deleted": {"$ne": True}})
        meal_label = next(m["label"] for m in MEALS if m["code"] == row["meal"])
        diet_label = next(d["label"] for d in MATRIX_DIETS if d["code"] == row["diet"])
        patch = {
            "matrix_key": key,
            "day_code": row["day_code"],
            "day_label": row["day_label"],
            "menu_date": row["menu_date"],
            "category": row["meal"],
            "diet": row["diet"],
            "name": row["name"],
            "description": row.get("description") or "",
            "amount_paise": row["amount_paise"],
            "amount_label": f"₹{row['amount_paise'] // 100}",
            "active": row["active"],
            "updated_at": iso(),
            "is_deleted": False,
            "source": "matrix_excel",
        }
        if row.get("image_url"):
            patch["image_url"] = row["image_url"]
        elif not existing:
            patch["image_url"] = ""

        if existing:
            # Don't wipe image unless Excel provided a URL
            if "image_url" not in patch:
                pass
            await db.food_menu_items.update_one({"id": existing["id"]}, {"$set": patch})
            updated += 1
        else:
            from food_poll_catalog import DAYS as _DAYS
            day = next((d for d in _DAYS if d["code"] == row["day_code"]), {})
            await db.food_menu_items.insert_one({
                "id": new_id("fitem"),
                **patch,
                "image_url": patch.get("image_url") or "",
                "sort_order": int(day.get("order") or 0) * 100,
                "created_at": iso(),
                "created_by": user.get("user_id") or "",
            })
            created += 1

    await audit(
        "food.menu_matrix.import",
        actor=user,
        entity_type="food_menu_matrix",
        entity_id="import",
        after={"created": created, "updated": updated, "skipped": skipped, "errors": len(errors)},
        request=request,
    )
    return {
        "ok": True,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:20],
        "message": f"Imported matrix: {created} created, {updated} updated"
        + (f", {skipped} skipped" if skipped else ""),
    }


@router.post("/admin/food-menu-matrix/upload-overall")
async def admin_food_matrix_upload_overall(
    file: UploadFile = File(...),
    request: Request = None,
    user: dict = Depends(require("households:write", "ops:manage", "receipts:manage", "settings:manage")),
):
    """Upload overall food menu (Word/Excel/PDF/image). Stores file + fills matrix when parseable."""
    from food_menu_matrix import parse_overall_menu_file, DAYS as MATRIX_DAYS
    from storage import save_document

    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="Empty file.")
    if len(raw) > 20 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large (max 20 MB).")

    filename = file.filename or "food-menu"
    try:
        parsed = parse_overall_menu_file(raw, filename)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read that menu file.")

    ctype = getattr(file, "content_type", None) or "application/octet-stream"
    doc = await save_document(
        data=raw,
        filename=filename,
        doc_type="food_menu_overall",
        linked_type="food_subscription",
        linked_id="overall_menu",
        uploaded_by=user.get("user_id") or "admin",
        content_type=ctype,
    )
    overall_url = doc.get("storage_path") or ""

    s = await get_settings()
    food = dict(s.get("food_subscription") or {})
    food["overall_menu_url"] = overall_url
    food["overall_menu_filename"] = filename[:200]
    food["overall_menu_uploaded_at"] = iso()
    food["overall_menu_kind"] = parsed.get("kind") or "document"
    await db.application_settings.update_one({}, {"$set": {"food_subscription": food}}, upsert=True)

    created = updated = skipped = filled = 0
    errors = []
    rows = parsed.get("rows") or []
    apply_all = parsed.get("kind") == "matrix_excel"

    for row in rows:
        if row.get("error"):
            errors.append({"row": row.get("excel_row"), "error": row["error"]})
            skipped += 1
            continue
        if not apply_all and not row.get("filled"):
            continue
        filled += 1
        key = row["matrix_key"]
        existing = await db.food_menu_items.find_one({"matrix_key": key, "is_deleted": {"$ne": True}})
        if apply_all or not existing:
            amount_paise = int(row["amount_paise"])
        else:
            amount_paise = int(existing.get("amount_paise") or row["amount_paise"])
        description = (row.get("description") or "").strip()
        if not description and existing:
            description = existing.get("description") or ""
        patch = {
            "matrix_key": key,
            "day_code": row["day_code"],
            "day_label": row["day_label"],
            "menu_date": row["menu_date"],
            "category": row["meal"],
            "diet": row["diet"],
            "name": row["name"],
            "description": description,
            "amount_paise": amount_paise,
            "amount_label": f"₹{amount_paise // 100}",
            "active": bool(row.get("active", True)),
            "updated_at": iso(),
            "is_deleted": False,
            "source": "overall_menu_upload",
        }
        if apply_all and row.get("image_url"):
            patch["image_url"] = row["image_url"]

        if existing:
            await db.food_menu_items.update_one({"id": existing["id"]}, {"$set": patch})
            updated += 1
        else:
            day = next((d for d in MATRIX_DAYS if d["code"] == row["day_code"]), {})
            meal = row["meal"]
            diet = row["diet"]
            await db.food_menu_items.insert_one({
                "id": new_id("fitem"),
                **patch,
                "image_url": patch.get("image_url") or "",
                "sort_order": int(day.get("order") or 0) * 100
                    + _MEAL_SORT_TENS.get(meal, 40)
                    + (1 if diet == "veg" else 2),
                "created_at": iso(),
                "created_by": user.get("user_id") or "",
            })
            created += 1

    note = parsed.get("note") or "Overall menu uploaded"
    message = note
    if created or updated:
        message = f"{note}. Matrix: {created} created, {updated} updated"
    elif parsed.get("kind") == "document":
        message = f"{note} Residents can open it on the Food page."

    await audit(
        "food.menu_matrix.upload_overall",
        actor=user,
        entity_type="food_menu_matrix",
        entity_id="overall",
        after={
            "filename": filename,
            "url": overall_url,
            "kind": parsed.get("kind"),
            "created": created,
            "updated": updated,
            "filled": filled,
        },
        request=request,
    )
    return {
        "ok": True,
        "message": message,
        "kind": parsed.get("kind"),
        "overall_menu_url": overall_url,
        "overall_menu_filename": filename,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "filled": filled,
        "errors": errors[:20],
        "note": note,
    }


@router.put("/admin/food-subscriptions/page")
async def admin_set_food_page(
    body: dict = Body(...),
    request: Request = None,
    user: dict = Depends(require("households:write", "ops:manage", "receipts:manage")),
):
    """One-click open / coming-soon for the public Food page."""
    s = await get_settings()
    food = dict(s.get("food_subscription") or {})
    if not food:
        raise HTTPException(status_code=503, detail="Food subscription is not configured yet.")
    if "page_enabled" not in body and "community_closed" not in body:
        raise HTTPException(status_code=400, detail="page_enabled or community_closed is required.")
    before = {"page_enabled": _food_page_enabled(food), "community_closed": _community_closed(food)}
    if "page_enabled" in body:
        food["page_enabled"] = bool(body.get("page_enabled"))
    if "community_closed" in body:
        food["community_closed"] = bool(body.get("community_closed"))
    if "closed_message" in body:
        food["closed_message"] = str(body.get("closed_message") or "")[:300]
    if "coming_soon_message" in body:
        food["coming_soon_message"] = str(body.get("coming_soon_message") or "")[:300]
    await db.application_settings.update_one({}, {"$set": {"food_subscription": food}}, upsert=True)
    await audit(
        "food.page.update",
        actor=user,
        entity_type="food_subscription",
        entity_id="page",
        before=before,
        after={"page_enabled": _food_page_enabled(food), "community_closed": _community_closed(food)},
        request=request,
    )
    return {
        "ok": True,
        "page_enabled": _food_page_enabled(food),
        "community_closed": _community_closed(food),
        "closed_message": _closed_message(food),
        "coming_soon_message": (food.get("coming_soon_message") or "").strip(),
    }


@router.put("/admin/food-subscriptions/prices")
async def admin_update_food_prices(
    body: dict = Body(...),
    request: Request = None,
    user: dict = Depends(require("households:write", "ops:manage", "receipts:manage")),
):
    """Update breakfast/lunch/dinner prices (rupees) and optionally open payment."""
    s = await get_settings()
    food = dict(s.get("food_subscription") or {})
    if not food:
        raise HTTPException(status_code=503, detail="Food subscription is not configured yet.")

    prices_in = body.get("prices") or {}
    # Accept {breakfast: 60, lunch: 300, dinner: 300} in rupees
    paise_map = {}
    for code in ("breakfast", "lunch", "dinner"):
        if code not in prices_in:
            continue
        try:
            rupees = float(prices_in[code])
        except Exception:
            raise HTTPException(status_code=400, detail=f"Invalid price for {code}")
        if rupees < 0:
            raise HTTPException(status_code=400, detail=f"Price for {code} cannot be negative")
        paise_map[code] = int(round(rupees * 100))

    days = []
    for day in food.get("days") or []:
        meals = []
        for meal in day.get("meals") or []:
            m = dict(meal)
            code = m.get("code")
            if code in paise_map:
                m["amount_paise"] = paise_map[code]
                m["amount_label"] = f"₹{paise_map[code] // 100}"
                m["amount_status"] = "fixed"
            meals.append(m)
        d = dict(day)
        d["meals"] = meals
        days.append(d)
    food["days"] = days

    if "page_enabled" in body:
        food["page_enabled"] = bool(body.get("page_enabled"))
    if "coming_soon_message" in body:
        food["coming_soon_message"] = str(body.get("coming_soon_message") or "")[:300]
    if "payment_enabled" in body:
        food["payment_enabled"] = bool(body.get("payment_enabled"))
    if "payment_note" in body:
        food["payment_note"] = str(body.get("payment_note") or "")[:300]
    if food.get("payment_enabled") and not food.get("payment_note"):
        food["payment_note"] = "Pay via UPI QR after selecting meals, then upload your payment screenshot."

    await db.application_settings.update_one({}, {"$set": {"food_subscription": food}}, upsert=True)
    await audit(
        "food.prices.update",
        actor=user,
        entity_type="food_subscription",
        entity_id="menu",
        after={"prices": paise_map, "payment_enabled": food.get("payment_enabled")},
        request=request,
    )
    food = _apply_default_prices(food)
    return {
        "ok": True,
        "page_enabled": _food_page_enabled(food),
        "coming_soon_message": (food.get("coming_soon_message") or "").strip(),
        "payment_enabled": bool(food.get("payment_enabled")),
        "payment_note": food.get("payment_note") or "",
        "prices": list(_meal_price_map(food).values()),
    }


@router.get("/admin/food-subscriptions")
async def admin_list_food(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    items = await db.food_subscriptions.find({}, NO_ID).sort("created_at", -1).to_list(2000)
    for it in items:
        it["voucher_token"] = _food_voucher_token(it["id"])
        settlement = await _order_settlement(it)
        it["paid_in_full"] = settlement["settled"]
        it["bank_verified"] = settlement["bank_verified"]
        it["ref"] = it["id"][-6:].upper()
        it["amount_due_paise"] = max(0, int(it.get("total_amount_paise") or 0) - int(settlement["paid"] or 0))
    s = await get_settings()
    food = s.get("food_subscription") or {}
    return {
        "items": items,
        "page_enabled": _food_page_enabled(food),
        "community_closed": _community_closed(food),
        "payment_enabled": bool(food.get("payment_enabled")),
    }


_CONFIRMED_FOOD_STATUSES = ("paid", "complimentary")


def _blank_counts() -> dict:
    return {"heads": 0, "kids": 0, "free": 0, "paid_plates": 0, "amount_paise": 0, "free_value_paise": 0, "flats": 0}


@router.get("/admin/food-reports")
async def admin_food_reports(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    """Kitchen headcount per day × meal × diet, one consolidated row per flat, and follow-up lists.

    Every live order counts (cancelled and expired ones don't). An order is "confirmed" once it is paid
    or fully complimentary; everything else is "awaiting" payment. Heads include the ₹0 kid plates.
    """
    from docs import receipt_is_bank_verified
    orders = await db.food_subscriptions.find(
        {"payment_status": {"$nin": _RELEASED_PAYMENT_STATUSES}}, NO_ID,
    ).sort("created_at", 1).to_list(5000)
    receipts_by: dict = {}
    async for r in db.receipts.find(
        {"food_subscription_id": {"$in": [o["id"] for o in orders]}, "status": "issued"},
        {"_id": 0, "food_subscription_id": 1, "receipt_no": 1, "total_amount": 1, "issued_at": 1,
         "bank_verified": 1, "issuance_source": 1, "masked_ref": 1},
    ).sort("issued_at", 1):
        receipts_by.setdefault(r["food_subscription_id"], []).append({
            "receipt_no": r.get("receipt_no"), "amount_paise": int(r.get("total_amount") or 0),
            "issued_at": r.get("issued_at"), "ref": r.get("masked_ref") or "",
            "bank_verified": receipt_is_bank_verified(r),
        })

    def stage_of(o):
        return "confirmed" if o.get("payment_status") in _CONFIRMED_FOOD_STATUSES else "awaiting"

    def flat_key(o):
        return o.get("flat_id") or f"{o.get('tower_name')}|{o.get('flat_number')}"

    # An unpaid order with exactly the same meals as another live order of the flat is usually a retry:
    # list it, but keep it out of plates, billed and due.
    duplicate_of: dict = {}
    first_seen: dict = {}
    for o in sorted(orders, key=lambda o: (stage_of(o) != "confirmed", o.get("created_at") or "")):
        sig = (flat_key(o), tuple(sorted(
            (s.get("item_id") or f"{s.get('day_code')}|{s.get('meal_code')}",
             int(s.get("quantity") or 0), int(s.get("kid_qty") or 0))
            for s in o.get("selections") or []
        )))
        if sig in first_seen and stage_of(o) == "awaiting":
            duplicate_of[o["id"]] = first_seen[sig]
        else:
            first_seen.setdefault(sig, o["id"][-6:].upper())

    meals: dict = {}
    flats: dict = {}
    for o in orders:
        stage = stage_of(o)
        receipts = receipts_by.get(o["id"], [])
        total = int(o.get("total_amount_paise") or 0)
        paid = sum(r["amount_paise"] for r in receipts)
        fkey = flat_key(o)
        flat = flats.setdefault(fkey, {
            "flat_id": o.get("flat_id") or "", "tower_id": o.get("tower_id") or "",
            "tower_name": o.get("tower_name") or "", "flat_number": o.get("flat_number") or "",
            "names": [], "mobiles": [], "orders": [], "lines": {},
            "billed_paise": 0, "paid_paise": 0, "due_paise": 0, "duplicate_due_paise": 0,
            "heads_confirmed": 0, "heads_awaiting": 0, "kids": 0, "complimentary": {},
            "subscribed": False,
        })
        for field, value in (("names", (o.get("name") or "").strip()), ("mobiles", o.get("mobile") or "")):
            if value and value not in flat[field]:
                flat[field].append(value)
        flat["subscribed"] = flat["subscribed"] or bool(o.get("flat_subscribed"))
        lines = _voucher_lines(o)
        due = max(0, total - paid) if stage == "awaiting" else 0
        flat["orders"].append({
            "id": o["id"], "ref": o["id"][-6:].upper(), "stage": stage, "payment_status": o.get("payment_status"),
            "name": o.get("name") or "", "created_at": o.get("created_at"),
            "total_paise": total, "paid_paise": paid, "due_paise": due,
            "receipts": receipts, "voucher_no": o.get("voucher_no") or "", "coupon_no": o.get("coupon_no") or "",
            "coupons_given": bool(o.get("coupon_id")), "whatsapp_posted": bool(o.get("whatsapp_posted_at")),
            "plates": sum(ln["heads"] for ln in lines), "duplicate_of": duplicate_of.get(o["id"], ""),
        })
        if o["id"] in duplicate_of:
            flat["duplicate_due_paise"] += due
            continue
        flat["billed_paise"] += total
        flat["paid_paise"] += paid
        flat["due_paise"] += due
        for k, v in (o.get("complimentary") or {}).items():
            flat["complimentary"][k] = flat["complimentary"].get(k, 0) + int(v or 0)
        for ln in lines:
            key = f"{ln['day_code']}|{ln['meal_code']}|{ln['diet']}"
            cell = meals.setdefault(key, {
                "day_code": ln["day_code"], "day_label": ln["day_label"], "date": ln["date"],
                "weekday": ln["weekday"], "day_order": ln["day_order"],
                "meal_code": ln["meal_code"], "meal_label": ln["meal_label"],
                "diet": ln["diet"], "diet_label": ln["diet_label"], "unit_paise": int(ln["unit_paise"] or 0),
                "confirmed": _blank_counts(), "awaiting": _blank_counts(), "_flats": {"confirmed": set(), "awaiting": set()},
            })
            c = cell[stage]
            c["heads"] += ln["heads"]
            c["kids"] += ln["kids"]
            c["free"] += ln["free"]
            c["paid_plates"] += ln["paid"]
            c["amount_paise"] += int(ln["amount_paise"] or 0)
            c["free_value_paise"] += ln["free"] * int(ln["unit_paise"] or 0)
            cell["_flats"][stage].add(fkey)
            fl = flat["lines"].setdefault(key, {
                "day_code": ln["day_code"], "day_label": ln["day_label"], "day_order": ln["day_order"],
                "meal_code": ln["meal_code"], "meal_label": ln["meal_label"],
                "diet": ln["diet"], "diet_label": ln["diet_label"],
                "confirmed": 0, "awaiting": 0, "kids": 0, "free": 0,
            })
            fl[stage] += ln["heads"]
            fl["kids"] += ln["kids"]
            fl["free"] += ln["free"]
            flat["heads_confirmed" if stage == "confirmed" else "heads_awaiting"] += ln["heads"]
            flat["kids"] += ln["kids"]

    meal_rows = []
    for cell in meals.values():
        for stage, keys in cell.pop("_flats").items():
            cell[stage]["flats"] = len(keys)
        meal_rows.append(cell)
    meal_rows.sort(key=lambda c: (c["day_order"], _VOUCHER_MEAL_ORDER.get(c["meal_code"], 9), c["diet"] != "veg"))

    flat_rows = []
    for flat in flats.values():
        flat["lines"] = sorted(flat["lines"].values(),
                               key=lambda x: (x["day_order"], _VOUCHER_MEAL_ORDER.get(x["meal_code"], 9), x["diet"] != "veg"))
        live = flat["orders"]
        flat["order_count"] = len(live)
        flat["awaiting_orders"] = sum(1 for x in live if x["stage"] == "awaiting" and not x["duplicate_of"])
        flat["unverified_receipts"] = sum(1 for x in live for r in x["receipts"] if not r["bank_verified"])
        flat["coupons_pending"] = sum(1 for x in live if x["stage"] == "confirmed" and not x["coupons_given"])
        flat["free_used"] = sum(flat["complimentary"].values())
        flat_rows.append(flat)
    flat_rows.sort(key=lambda f: (f["tower_id"] or "", f["flat_number"]))

    # Flats that paid the Puja subscription (complimentary meals unlocked) but have no live food order.
    ordered = {f["flat_id"] for f in flat_rows if f["flat_id"]}
    cycle = await get_active_cycle_id()
    sub_house = {}
    async for r in db.receipts.find(
        {"kind": "subscription", "status": "issued", "cycle_id": cycle},
        {"_id": 0, "household_id": 1, "payer_name": 1, "payer_mobile": 1},
    ).sort("issued_at", -1):
        sub_house.setdefault(r.get("household_id"), r)
    not_ordered = []
    async for h in db.households.find(
        {"id": {"$in": [k for k in sub_house if k]}, "is_deleted": {"$ne": True}},
        {"_id": 0, "id": 1, "flat_id": 1, "tower_id": 1, "tower_name": 1, "flat_number": 1,
         "primary_name": 1, "primary_mobile": 1},
    ):
        if not h.get("flat_id") or h["flat_id"] in ordered:
            continue
        ordered.add(h["flat_id"])
        r = sub_house[h["id"]]
        contact = _real_contact(r.get("payer_name"), r.get("payer_mobile")) or \
            _real_contact(h.get("primary_name"), h.get("primary_mobile"))
        not_ordered.append({
            "flat_id": h["flat_id"], "tower_id": h.get("tower_id") or "", "tower_name": h.get("tower_name") or "",
            "flat_number": h.get("flat_number") or "", "name": contact.get("name") or "",
            "mobile": contact.get("mobile") or "",
        })
    not_ordered.sort(key=lambda f: (f["tower_id"], f["flat_number"]))

    vendor_plates = vendor_value = vendor_entries = 0
    async for v in db.vendor_meals.find({"status": "active"}, NO_ID):
        vendor_entries += 1
        for ln in _voucher_lines(v):
            key = f"{ln['day_code']}|{ln['meal_code']}|{ln['diet']}"
            cell = next((c for c in meal_rows if f"{c['day_code']}|{c['meal_code']}|{c['diet']}" == key), None)
            if cell is None:
                cell = {
                    "day_code": ln["day_code"], "day_label": ln["day_label"], "date": ln["date"],
                    "weekday": ln["weekday"], "day_order": ln["day_order"],
                    "meal_code": ln["meal_code"], "meal_label": ln["meal_label"],
                    "diet": ln["diet"], "diet_label": ln["diet_label"], "unit_paise": int(ln["unit_paise"] or 0),
                    "confirmed": _blank_counts(), "awaiting": _blank_counts(),
                }
                meal_rows.append(cell)
            c = cell.setdefault("vendor", _blank_counts())
            c["heads"] += ln["heads"]
            c["amount_paise"] += int(ln["amount_paise"] or 0)
            vendor_plates += ln["heads"]
            vendor_value += int(ln["amount_paise"] or 0)
    for cell in meal_rows:
        cell.setdefault("vendor", _blank_counts())
    meal_rows.sort(key=lambda c: (c["day_order"], _VOUCHER_MEAL_ORDER.get(c["meal_code"], 9), c["diet"] != "veg"))

    from pujo_menu_2026 import COMPLIMENTARY_RULES
    return {
        "generated_at": iso(),
        "meals": meal_rows,
        "flats": flat_rows,
        "subscribed_not_ordered": not_ordered,
        "complimentary_rules": [{"key": r["key"], "free": r["free"], "label": r.get("label") or r["key"],
                                 "day_code": r.get("day_code") or ""} for r in COMPLIMENTARY_RULES],
        "totals": {
            "orders": len(orders),
            "flats": len(flat_rows),
            "billed_paise": sum(f["billed_paise"] for f in flat_rows),
            "paid_paise": sum(f["paid_paise"] for f in flat_rows),
            "due_paise": sum(f["due_paise"] for f in flat_rows),
            "duplicate_due_paise": sum(f["duplicate_due_paise"] for f in flat_rows),
            "vendor_entries": vendor_entries,
            "vendor_plates": vendor_plates,
            "vendor_value_paise": vendor_value,
        },
    }


def _cache_food_revenue_files(pdf: bytes, png: bytes) -> None:
    """Keep /reports/… copies fresh so committee can open them without the API blob path."""
    from pathlib import Path
    roots = [
        Path(__file__).resolve().parent.parent / "frontend" / "build" / "reports",
        Path(__file__).resolve().parent.parent / "frontend" / "public" / "reports",
        Path("/var/www/one10events/build/reports"),
    ]
    for root in roots:
        try:
            root.mkdir(parents=True, exist_ok=True)
            (root / "one10_food_revenue.pdf").write_bytes(pdf)
            (root / "one10_food_revenue.png").write_bytes(png)
        except OSError:
            pass


async def _food_revenue_payload():
    from food_revenue_report import (
        build_food_revenue_report, food_revenue_report_pdf, food_revenue_report_png,
    )
    report = await build_food_revenue_report()
    pdf = food_revenue_report_pdf(report)
    png = food_revenue_report_png(report)
    _cache_food_revenue_files(pdf, png)
    return report, pdf, png


def _revenue_filename(ext: str) -> str:
    """Safe ASCII filename — colons/spaces in Content-Disposition break downloads in some browsers."""
    from util import to_ist, now_utc
    stamp = to_ist(now_utc()).strftime("%Y%m%d_%H%M")
    return f"one10_food_revenue_{stamp}.{ext}"


@router.get("/admin/food-reports/revenue-pdf")
@router.get("/admin/food-reports/revenue.pdf")
async def admin_food_revenue_pdf(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    """Consolidated food coupon revenue PDF — veg breakfast packets excluded (already paid)."""
    _report, pdf, _png = await _food_revenue_payload()
    name = _revenue_filename("pdf")
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{name}"',
            "Cache-Control": "no-store",
        },
    )


@router.get("/admin/food-reports/revenue-png")
@router.get("/admin/food-reports/revenue.png")
async def admin_food_revenue_png(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    """Same consolidated revenue summary as a shareable PNG (packets excluded)."""
    _report, _pdf, png = await _food_revenue_payload()
    name = _revenue_filename("png")
    return Response(
        content=png,
        media_type="image/png",
        headers={
            "Content-Disposition": f'attachment; filename="{name}"',
            "Cache-Control": "no-store",
        },
    )


@router.get("/admin/food-reports/cost-revenue-pdf")
async def admin_food_cost_revenue_pdf(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    """P&L PDF: revenue vs caterer cost with complimentary plates costed (paid to vendor)."""
    from food_cost_revenue_report import generate_food_cost_revenue_pdf, publish_pdf
    pdf, _revenue, _expense = await generate_food_cost_revenue_pdf()
    publish_pdf(pdf)
    name = _revenue_filename("pdf").replace("food_revenue", "food_cost_revenue")
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{name}"',
            "Cache-Control": "no-store",
        },
    )


# Plates the committee arranges for vendors (dhakis, priests, decorators, security…). Nobody pays for them,
# so they are a cost to the committee: every entry keeps who added it, and entries are cancelled, never deleted.
VENDOR_MEAL_TYPES = {
    "dhaki": "Dhaki / musicians",
    "priest": "Priest / purohit",
    "decorator": "Decorator / pandal",
    "light_sound": "Light & sound",
    "security": "Security",
    "housekeeping": "Housekeeping / cleaning",
    "caterer_staff": "Caterer staff",
    "other": "Other",
}
_VENDOR_MEAL_READ = ("households:read", "ops:read", "receipts:read", "reports:read", "expense:create")
_VENDOR_MEAL_WRITE = ("households:write", "ops:manage", "receipts:manage", "expense:create")


def _vendor_meal_public(doc: dict) -> dict:
    lines = _voucher_lines(doc)
    return {
        **{k: doc.get(k) for k in (
            "id", "ref", "vendor_name", "vendor_type", "contact_mobile", "purpose", "status",
            "plates", "value_paise", "created_at", "created_by_name",
            "cancelled_at", "cancelled_by_name", "cancel_reason",
        )},
        "vendor_type_label": VENDOR_MEAL_TYPES.get(doc.get("vendor_type") or "") or doc.get("vendor_type") or "",
        "lines": lines,
        "days": _voucher_days(lines),
    }


@router.get("/admin/vendor-meals")
async def admin_list_vendor_meals(user: dict = Depends(require(*_VENDOR_MEAL_READ))):
    docs = await db.vendor_meals.find({}, NO_ID).sort("created_at", -1).to_list(2000)
    active = [d for d in docs if d.get("status") == "active"]
    return {
        "items": [_vendor_meal_public(d) for d in docs],
        "types": [{"value": k, "label": v} for k, v in VENDOR_MEAL_TYPES.items()],
        "can_edit": any(p in user_permissions(user) for p in _VENDOR_MEAL_WRITE),
        "totals": {
            "entries": len(active),
            "plates": sum(int(d.get("plates") or 0) for d in active),
            "value_paise": sum(int(d.get("value_paise") or 0) for d in active),
            "cancelled": len(docs) - len(active),
        },
    }


@router.post("/admin/vendor-meals")
async def admin_create_vendor_meal(body: dict = Body(...), request: Request = None,
                                   user: dict = Depends(require(*_VENDOR_MEAL_WRITE))):
    vendor_name = (body.get("vendor_name") or "").strip()[:120]
    vendor_type = (body.get("vendor_type") or "").strip()
    purpose = (body.get("purpose") or "").strip()[:500]
    mobile = re.sub(r"\D", "", str(body.get("contact_mobile") or ""))[-10:]
    if not vendor_name:
        raise HTTPException(status_code=400, detail="Enter the vendor's name.")
    if vendor_type not in VENDOR_MEAL_TYPES:
        raise HTTPException(status_code=400, detail="Choose the vendor type.")
    if len(purpose) < 3:
        raise HTTPException(status_code=400, detail="Say why these meals are given (shown in the audit).")
    if mobile and not re.fullmatch(r"[6-9]\d{9}", mobile):
        raise HTTPException(status_code=400, detail="Enter a valid 10-digit mobile, or leave it blank.")

    items = await db.food_menu_items.find({"is_deleted": {"$ne": True}}, NO_ID).to_list(500)
    items_by_id = {i["id"]: {**i, "active": True} for i in items}
    s = await get_settings()
    selections = _normalize_selections(_apply_default_prices(s.get("food_subscription") or {}),
                                       [r for r in body.get("selections") or [] if isinstance(r, dict) and r.get("item_id")],
                                       items_by_id=items_by_id)
    if not selections:
        raise HTTPException(status_code=400, detail="Add at least one meal.")
    for sel in selections:
        sel["kid_qty"] = 0
        sel["complimentary_qty"] = 0
        sel["line_total_paise"] = int(sel.get("amount_paise") or 0) * int(sel.get("quantity") or 0)

    _, ref = await next_formatted("vendor_meal", "VM26", 4)
    doc = {
        "id": new_id("vmeal"),
        "ref": ref,
        "cycle_id": await get_active_cycle_id(),
        "vendor_name": vendor_name,
        "vendor_type": vendor_type,
        "contact_mobile": mobile,
        "purpose": purpose,
        "selections": selections,
        "plates": sum(int(x.get("quantity") or 0) for x in selections),
        "value_paise": sum(int(x["line_total_paise"]) for x in selections),
        "status": "active",
        "created_at": iso(),
        "created_by": user.get("user_id") or "",
        "created_by_name": user.get("name") or user.get("login_id") or "",
    }
    await db.vendor_meals.insert_one(dict(doc))
    await audit("food.vendor_meal.create", actor=user, entity_type="vendor_meal", entity_id=doc["id"],
                reason=purpose,
                after={"ref": ref, "vendor_name": vendor_name, "vendor_type": vendor_type,
                       "plates": doc["plates"], "value_paise": doc["value_paise"],
                       "items": [{"name": x.get("name"), "menu_date": x.get("menu_date"),
                                  "quantity": x.get("quantity"), "amount_paise": x.get("amount_paise")}
                                 for x in selections]},
                request=request)
    return {"ok": True, "item": _vendor_meal_public(doc)}


@router.post("/admin/vendor-meals/{meal_id}/cancel")
async def admin_cancel_vendor_meal(meal_id: str, body: dict = Body(default={}), request: Request = None,
                                   user: dict = Depends(require(*_VENDOR_MEAL_WRITE))):
    reason = (body.get("reason") or "").strip()[:500]
    if len(reason) < 3:
        raise HTTPException(status_code=400, detail="Give a reason for cancelling.")
    doc = await db.vendor_meals.find_one({"id": meal_id}, NO_ID)
    if not doc:
        raise HTTPException(status_code=404, detail="Vendor meal entry not found.")
    if doc.get("status") != "active":
        raise HTTPException(status_code=409, detail="This entry is already cancelled.")
    update = {
        "status": "cancelled",
        "cancelled_at": iso(),
        "cancelled_by": user.get("user_id") or "",
        "cancelled_by_name": user.get("name") or user.get("login_id") or "",
        "cancel_reason": reason,
    }
    await db.vendor_meals.update_one({"id": meal_id, "status": "active"}, {"$set": update})
    await audit("food.vendor_meal.cancel", actor=user, entity_type="vendor_meal", entity_id=meal_id,
                reason=reason, before={"status": "active", "plates": doc.get("plates"), "value_paise": doc.get("value_paise")},
                after={"status": "cancelled"}, request=request)
    return {"ok": True, "item": _vendor_meal_public({**doc, **update})}


_BASE_TEMPLATE_COLS = [
    "id",
    "name",
    "mobile",
    "tower_name",
    "flat_number",
    "family_members",
    "notes",
    "payment_status",
    "status",
]
_TRUTHY = {"y", "yes", "1", "true", "t", "x", "✓", "✔"}


def _meal_slot_headers(food: dict) -> list[str]:
    food = _strip_evening_menu(food or {})
    headers = []
    for day in food.get("days") or []:
        dcode = (day.get("code") or "").strip()
        if not dcode:
            continue
        for meal in day.get("meals") or []:
            mcode = (meal.get("code") or "").strip()
            if mcode:
                headers.append(f"{dcode}_{mcode}")
    return headers


def _normalize_header(h: str) -> str:
    return re.sub(r"[^a-z0-9_]+", "_", (h or "").strip().lower()).strip("_")


def _is_truthy(val) -> bool:
    if val is None:
        return False
    s = str(val).strip().lower()
    return s in _TRUTHY


def _parse_spreadsheet(content: bytes, filename: str) -> list[list]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        return [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
    text = content.decode("utf-8-sig", errors="replace")
    return list(csv.reader(io.StringIO(text)))


def _row_dict(header: list[str], row: list) -> dict:
    out = {}
    for i, key in enumerate(header):
        if not key:
            continue
        out[key] = row[i] if i < len(row) else ""
    return out


def _selections_from_row(food: dict, row: dict, meal_headers: list[str]) -> list | None:
    """Return selections from Y/N meal columns, or None if file has no meal columns."""
    present = [h for h in meal_headers if h in row]
    if not present:
        # Also accept a compact "meals" column: day|meal;day|meal
        raw = str(row.get("meals") or "").strip()
        if not raw:
            return None
        keys = [p.strip() for p in re.split(r"[;,]", raw) if p.strip()]
        return _normalize_selections(food, keys)

    lines = []
    for h in present:
        raw = str(row.get(h) or "").strip()
        if not raw:
            continue
        if "_" not in h:
            continue
        day, meal = h.split("_", 1)
        # Numeric quantity (e.g. 3) or truthy Y/yes
        qty = None
        try:
            qty = int(float(raw))
        except (TypeError, ValueError):
            qty = 1 if _is_truthy(raw) else 0
        if qty and qty > 0:
            lines.append({"day_code": day, "meal_code": meal, "quantity": qty})
    return _normalize_selections(food, lines)


def _payment_status_from_cell(val: str) -> str:
    s = (val or "").strip().lower()
    if s in {"paid", "captured", "settled", "yes", "y", "1"}:
        return "paid"
    if s in {"unpaid", "not_open", "not open", "pending", "no", "n", "0"}:
        return "not_open"
    if s:
        return s.replace(" ", "_")
    return "not_open"


@router.get("/admin/food-subscriptions/template")
async def admin_food_template(
    include_data: bool = True,
    user: dict = Depends(require("households:read", "ops:read", "receipts:read")),
):
    """CSV template for offline capture. Fill meal columns with Y/N, then upload."""
    s = await get_settings()
    food = _strip_evening_menu(s.get("food_subscription") or {})
    meal_headers = _meal_slot_headers(food)
    headers = _BASE_TEMPLATE_COLS + meal_headers

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)

    # Instruction / example blank row guidance as comment-like first data hint
    if include_data:
        items = await db.food_subscriptions.find({}, NO_ID).sort("created_at", -1).to_list(5000)
        for item in items:
            chosen = {
                f"{(sel.get('day_code') or '')}_{(sel.get('meal_code') or '')}"
                for sel in (item.get("selections") or [])
            }
            row = [
                item.get("id") or "",
                item.get("name") or "",
                item.get("mobile") or "",
                item.get("tower_name") or "",
                item.get("flat_number") or "",
                item.get("family_members") or 1,
                item.get("notes") or "",
                item.get("payment_status") or "not_open",
                item.get("status") or "registered",
            ]
            row.extend(["Y" if h in chosen else "" for h in meal_headers])
            writer.writerow(row)
    else:
        # One blank sample row so Excel opens with columns ready
        writer.writerow([""] * len(headers))

    data = buf.getvalue().encode("utf-8-sig")
    return StreamingResponse(
        io.BytesIO(data),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="food_subscription_template.csv"'},
    )


@router.post("/admin/food-subscriptions/import")
async def admin_food_import(
    file: UploadFile = File(...),
    request: Request = None,
    user: dict = Depends(require("households:write", "ops:manage", "receipts:manage")),
):
    """Upload filled template (CSV/XLSX) to create or update food subscriptions."""
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Empty file.")

    rows = _parse_spreadsheet(content, file.filename or "upload.csv")
    if not rows or len(rows) < 2:
        raise HTTPException(status_code=400, detail="File needs a header row and at least one data row.")

    header = [_normalize_header(str(h)) for h in rows[0]]
    if "name" not in header or "mobile" not in header:
        raise HTTPException(status_code=400, detail="Template must include name and mobile columns.")

    s = await get_settings()
    food = s.get("food_subscription") or {}
    if not food:
        raise HTTPException(status_code=503, detail="Food subscription is not configured yet.")
    meal_headers = _meal_slot_headers(food)
    cycle_id = await get_active_cycle_id()

    created = updated = skipped = 0
    errors = []

    for idx, raw in enumerate(rows[1:], start=2):
        if not any(str(x).strip() for x in (raw or [])):
            continue
        row = _row_dict(header, raw)
        name = str(row.get("name") or "").strip()
        mobile = "".join(c for c in str(row.get("mobile") or "") if c.isdigit())[-10:]
        if not name or len(mobile) != 10:
            skipped += 1
            errors.append({"row": idx, "error": "Name and valid 10-digit mobile required"})
            continue

        selections = _selections_from_row(food, row, meal_headers)
        sid = str(row.get("id") or "").strip()
        existing = None
        if sid:
            existing = await db.food_subscriptions.find_one({"id": sid})
        if not existing:
            existing = await db.food_subscriptions.find_one({"mobile": mobile}, sort=[("created_at", -1)])
        if existing and (existing.get("voucher_no") or existing.get("coupon_id") or existing.get("intent_id")
                         or existing.get("payment_status") in _VOUCHER_ELIGIBLE):
            skipped += 1
            errors.append({"row": idx, "error": "Locked: online, paid or vouchered orders can't be changed by upload"})
            continue

        family_members = 1
        try:
            family_members = max(1, int(float(str(row.get("family_members") or 1).strip() or 1)))
        except Exception:
            family_members = 1

        payment_status = _payment_status_from_cell(str(row.get("payment_status") or ""))
        status = (str(row.get("status") or "").strip() or (existing or {}).get("status") or "registered")
        notes = str(row.get("notes") or "").strip()[:500]
        tower_name = str(row.get("tower_name") or "").strip()
        flat_number = str(row.get("flat_number") or "").strip()

        patch = {
            "name": name,
            "mobile": mobile,
            "tower_name": tower_name,
            "flat_number": flat_number,
            "family_members": family_members,
            "notes": notes,
            "payment_status": payment_status,
            "status": status,
            "updated_at": iso(),
        }
        if selections is not None:
            if not selections:
                skipped += 1
                errors.append({"row": idx, "error": "Select at least one meal (put Y in meal columns)"})
                continue
            patch["selections"] = selections
            patch["selection_count"] = len(selections)

        if existing:
            await db.food_subscriptions.update_one({"id": existing["id"]}, {"$set": patch})
            updated += 1
            await audit(
                "food.subscription.import_update",
                actor=user,
                entity_type="food_subscription",
                entity_id=existing["id"],
                after={"name": name, "mobile": mobile[-4:], "selection_count": patch.get("selection_count")},
                reason=f"Template upload row {idx}",
                request=request,
            )
        else:
            if selections is None or not selections:
                skipped += 1
                errors.append({"row": idx, "error": "New rows need at least one meal marked Y"})
                continue
            new_sid = new_id("food")
            doc = {
                "id": new_sid,
                "cycle_id": cycle_id,
                "tower_id": "",
                "flat_id": "",
                "amount_status": "TBC",
                "total_amount_paise": None,
                "coupon_id": None,
                "coupon_no": None,
                "created_at": iso(),
                "source": "admin_import",
                **patch,
            }
            await db.food_subscriptions.insert_one(doc)
            created += 1
            await audit(
                "food.subscription.import_create",
                actor=user,
                entity_type="food_subscription",
                entity_id=new_sid,
                after={"name": name, "mobile": mobile[-4:], "selection_count": len(selections)},
                reason=f"Template upload row {idx}",
                request=request,
            )

    return {
        "ok": True,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:50],
        "message": f"Import done — {created} created, {updated} updated, {skipped} skipped.",
    }


@router.get("/admin/food-subscriptions/{sid}")
async def admin_get_food(sid: str, user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    doc = await db.food_subscriptions.find_one({"id": sid}, NO_ID)
    if not doc:
        raise HTTPException(status_code=404, detail="Food subscription not found.")
    coupon = None
    if doc.get("coupon_id"):
        coupon = await db.food_coupons.find_one({"id": doc["coupon_id"]}, NO_ID)
    trail = await db.audit_events.find(
        {"entity_type": "food_subscription", "entity_id": sid}, NO_ID
    ).sort("seq", 1).to_list(200)
    if doc.get("coupon_id"):
        trail += await db.audit_events.find(
            {"entity_type": "food_coupon", "entity_id": doc["coupon_id"]}, NO_ID
        ).sort("seq", 1).to_list(200)
        trail.sort(key=lambda e: e.get("seq") or 0)
    return {"item": doc, "coupon": coupon, "trail": trail}


# =============================================================== FOOD VOUCHERS
# A voucher is the resident's proof of a fully paid (or fully complimentary) order; the committee
# checks it and issues the physical food coupons against it.
_VOUCHER_ELIGIBLE = ("paid", "complimentary")
_VOUCHER_MEAL_ORDER = {"breakfast": 1, "breakfast_packet": 2, "lunch": 3, "dinner": 4}
_VOUCHER_MEAL_LABEL = {
    "breakfast": "Breakfast", "breakfast_packet": "Breakfast packet", "lunch": "Lunch", "dinner": "Dinner",
}


def _mask_mobile(mobile: str) -> str:
    m = (mobile or "").strip()
    return f"{'•' * max(0, len(m) - 4)}{m[-4:]}" if m else ""


def _voucher_lines(doc: dict) -> list:
    day_by = {d["code"]: d for d in DAYS}
    lines = []
    for s in doc.get("selections") or []:
        day = s.get("item_day_code") or s.get("day_code") or ""
        d = day_by.get(day) or {}
        meal = s.get("category") or s.get("meal_code") or ""
        qty = max(1, int(s.get("quantity") or 1))
        free = min(qty, int(s.get("complimentary_qty") or 0))
        kids = int(s.get("kid_qty") or 0)
        unit = s.get("amount_paise")
        amount = s.get("line_total_paise")
        if amount is None and unit not in (None, ""):
            amount = int(unit) * (qty - free)
        lines.append({
            "day_code": day,
            "day_label": d.get("short_label") or s.get("day_label") or day,
            "date": d.get("date") or s.get("menu_date") or "",
            "weekday": d.get("weekday") or "",
            "day_order": d.get("order", 99),
            "meal_code": meal,
            "meal_label": _VOUCHER_MEAL_LABEL.get(meal) or s.get("category_label") or meal.title(),
            "item": s.get("name") or s.get("meal_label") or "",
            "diet": s.get("diet") or "",
            "diet_label": s.get("diet_label") or "",
            "heads": qty + kids,
            "free": free,
            "paid": qty - free,
            "kids": kids,
            "unit_paise": unit,
            "amount_paise": amount,
        })
    lines.sort(key=lambda x: (x["day_order"], _VOUCHER_MEAL_ORDER.get(x["meal_code"], 9), x["item"]))
    return lines


def _voucher_days(lines: list) -> list:
    days: list = []
    for ln in lines:
        day = next((d for d in days if d["day_code"] == ln["day_code"]), None)
        if not day:
            day = {"day_code": ln["day_code"], "day_label": ln["day_label"], "date": ln["date"],
                   "weekday": ln["weekday"], "meals": [], "heads": 0}
            days.append(day)
        meal = next((m for m in day["meals"] if m["meal_label"] == ln["meal_label"]), None)
        if not meal:
            meal = {"meal_label": ln["meal_label"], "heads": 0}
            day["meals"].append(meal)
        meal["heads"] += ln["heads"]
        day["heads"] += ln["heads"]
    return days


async def _order_settlement(doc: dict) -> dict:
    """Live check that an order is really settled: free with nothing to pay, or issued receipts cover the total.

    Re-checked on every voucher view, coupon issue and coupon scan, so a voided receipt or an edited
    status cannot keep a voucher or coupon alive.
    """
    from docs import receipt_is_bank_verified
    receipts = await db.receipts.find(
        {"food_subscription_id": doc["id"], "status": "issued"}, {"_id": 0},
    ).sort("issued_at", 1).to_list(50)
    total = doc.get("total_amount_paise")
    paid = sum(int(r.get("total_amount") or 0) for r in receipts)
    status = doc.get("payment_status")
    if status == "complimentary":
        settled = total == 0 and bool(doc.get("complimentary"))
    elif status == "paid":
        settled = total is not None and int(total) > 0 and paid >= int(total)
    else:
        settled = False
    return {
        "settled": settled,
        "receipts": receipts,
        "paid": paid,
        "total": total,
        "bank_verified": all(receipt_is_bank_verified(r) for r in receipts),
        "unverified": [r for r in receipts if not receipt_is_bank_verified(r)],
    }


async def _ensure_food_voucher(doc: dict, settlement: dict, actor=None, request: Request = None) -> dict:
    """Give a settled order its voucher number once; later calls return the same number."""
    if doc.get("voucher_no") or not settlement["settled"]:
        return doc
    s = await get_settings()
    prefix = (s.get("food_subscription") or {}).get("voucher_prefix") or "ONE10-FV26"
    _, voucher_no = await next_formatted("food_voucher", prefix, 5)
    updated = await db.food_subscriptions.find_one_and_update(
        {"id": doc["id"], "voucher_no": {"$in": [None, ""]}},
        {"$set": {"voucher_no": voucher_no, "voucher_issued_at": iso(), "updated_at": iso()}},
        projection=NO_ID, return_document=True,
    )
    if not updated:
        return await db.food_subscriptions.find_one({"id": doc["id"]}, NO_ID) or doc
    await audit("food.voucher.issue", actor=actor or "public", entity_type="food_subscription",
                entity_id=doc["id"], after={"voucher_no": voucher_no}, request=request)
    return updated


async def _food_voucher_payload(doc: dict, settlement: dict) -> dict:
    lines = _voucher_lines(doc)
    receipts = settlement["receipts"]
    coupon = None
    if doc.get("coupon_id"):
        c = await db.food_coupons.find_one({"id": doc["coupon_id"]}, NO_ID) or {}
        coupon = {"coupon_no": c.get("coupon_no") or doc.get("coupon_no"), "issued_at": c.get("created_at"),
                  "issued_by_name": c.get("issued_by_name") or "", "slips": c.get("slip_count") or 0}
    eligible = settlement["settled"]
    released = doc.get("payment_status") in _RELEASED_PAYMENT_STATUSES
    voided = (bool(doc.get("voucher_no")) or released) and not eligible
    total = settlement["total"]
    paid = settlement["paid"]
    return {
        "eligible": eligible,
        "voided": voided,
        "voucher_no": doc.get("voucher_no") if (eligible or voided) else None,
        "voucher_issued_at": doc.get("voucher_issued_at"),
        "voucher_token": _food_voucher_token(doc["id"]),
        "order_id": doc["id"],
        "ordered_at": doc.get("created_at"),
        "payment_status": doc.get("payment_status"),
        "status_label": ("Cancelled" if doc.get("payment_status") == "cancelled" else "Expired") if released else
        "Not valid" if voided else (
            "Complimentary" if doc.get("payment_status") == "complimentary" else (
                "Paid in full" if eligible else "Payment pending")),
        "bank_verified": settlement["bank_verified"],
        "name": doc.get("name") or "",
        "mobile_masked": _mask_mobile(doc.get("mobile")),
        "tower_name": doc.get("tower_name") or "",
        "flat_number": doc.get("flat_number") or "",
        "lines": lines,
        "days": _voucher_days(lines),
        "totals": {
            "heads": sum(ln["heads"] for ln in lines),
            "free": sum(ln["free"] for ln in lines),
            "paid": sum(ln["paid"] for ln in lines),
            "kids": sum(ln["kids"] for ln in lines),
        },
        "total_amount_paise": total,
        "amount_paid_paise": paid,
        "amount_due_paise": max(0, int(total or 0) - paid) if not eligible else 0,
        "receipts": [
            {"receipt_no": r.get("receipt_no"), "amount_paise": r.get("total_amount"),
             "issued_at": r.get("issued_at"), "is_partial": bool(r.get("is_partial"))}
            for r in receipts
        ],
        "coupon": coupon,
    }


async def _food_by_voucher_token(token: str) -> dict:
    fid = read_food_voucher_token(token)
    doc = await db.food_subscriptions.find_one({"id": fid}, NO_ID) if fid else None
    if not doc:
        raise HTTPException(status_code=404, detail="Food voucher not found.")
    return doc


@router.get("/food/voucher/{token}")
async def food_voucher(token: str, request: Request = None):
    doc = await _food_by_voucher_token(token)
    settlement = await _order_settlement(doc)
    doc = await _ensure_food_voucher(doc, settlement, request=request)
    return await _food_voucher_payload(doc, settlement)


@router.get("/food/voucher/{token}/pdf")
async def food_voucher_pdf_public(token: str, request: Request = None):
    from docs import food_voucher_pdf
    doc = await _food_by_voucher_token(token)
    settlement = await _order_settlement(doc)
    if not settlement["settled"]:
        raise HTTPException(status_code=409, detail="The food voucher is available once the full amount is paid.")
    doc = await _ensure_food_voucher(doc, settlement, request=request)
    payload = await _food_voucher_payload(doc, settlement)
    s = await get_settings()
    pdf = food_voucher_pdf(payload, s)
    return Response(content=pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="{payload["voucher_no"]}.pdf"'})


# ---------------------------------------------------------------- MY ORDERS (buyer self-service)
_OPEN_INTENT_STATUSES = ["payment_pending", "pending", "authorised", "partially_paid"]
_LIVE_SUBMISSION_STATUSES = ["submitted", "processing", "llm_ok", "needs_review"]


async def _payment_under_review(doc: dict) -> bool:
    if not doc.get("intent_id"):
        return False
    return bool(await db.upi_submissions.find_one(
        {"intent_id": doc["intent_id"], "status": {"$in": _LIVE_SUBMISSION_STATUSES},
         "receipt_id": {"$in": [None, ""]}},
        {"_id": 1},
    ))


async def _resident_change_problem(doc: dict) -> str:
    """Why a resident can't edit or cancel this order themselves; '' when they can (nothing paid or in flight)."""
    status = doc.get("payment_status")
    ask = " Please ask the committee."
    if status in _RELEASED_PAYMENT_STATUSES:
        return "This order is already cancelled or expired."
    if doc.get("coupon_id"):
        return "Coupons were already given for this order, so it can't be changed here." + ask
    if status not in ("pending", "not_open", "complimentary") or (
            doc.get("intent_id") and await db.receipts.find_one({"intent_id": doc["intent_id"]}, {"_id": 1})):
        return "A payment is recorded on this order, so it can't be changed here." + ask
    if doc.get("whatsapp_posted_at"):
        return "You posted a payment for this order on WhatsApp, so it can't be changed here." + ask
    if await _payment_under_review(doc):
        return "A payment screenshot for this order is being checked, so it can't be changed here." + ask
    return ""


async def _food_by_edit_token(token: str) -> dict:
    sid = read_food_edit_token((token or "").strip())
    doc = await db.food_subscriptions.find_one({"id": sid}, NO_ID) if sid else None
    if not doc:
        raise HTTPException(status_code=404, detail="This link has expired — open the order again from My orders.")
    return doc


@router.get("/food/my-orders/edit")
async def food_order_for_edit(token: str, request: Request):
    """An unpaid order's cart, so the resident can change it; checking out with replaces_token swaps it."""
    s = await get_settings()
    food = s.get("food_subscription") or {}
    if _community_closed(food) and not await _committee_orderer(request):
        raise HTTPException(status_code=409, detail=_closed_message(food))
    doc = await _food_by_edit_token(token)
    if problem := await _resident_change_problem(doc):
        raise HTTPException(status_code=409, detail=problem)
    return {
        "id": doc["id"], "ref": doc["id"][-6:].upper(), "edit_token": token,
        "tower_id": doc.get("tower_id"), "flat_id": doc.get("flat_id"),
        "tower_name": doc.get("tower_name"), "flat_number": doc.get("flat_number"),
        "name": doc.get("name") or "", "mobile": doc.get("mobile") or "",
        "selections": [
            {"item_id": s.get("item_id"), "day_code": s.get("day_code"), "meal_code": s.get("meal_code"),
             "quantity": int(s.get("quantity") or 0), "kid_qty": int(s.get("kid_qty") or 0)}
            for s in doc.get("selections") or []
        ],
        "kids": [{"birth_year": k.get("birth_year")} for k in doc.get("kids") or [] if k.get("birth_year")],
    }


@router.post("/food/my-orders/cancel")
async def food_cancel_own_order(body: dict = Body(...), request: Request = None):
    """Resident cancels their own order while nothing is paid or being checked; frees its complimentary meals."""
    doc = await _food_by_edit_token(body.get("edit_token"))
    async with _flat_guard(doc.get("flat_id") or doc["id"]):
        doc = await db.food_subscriptions.find_one({"id": doc["id"]}, NO_ID)
        if problem := await _resident_change_problem(doc):
            raise HTTPException(status_code=409, detail=problem)
        await _cancel_food_order(doc, actor="public", reason="Cancelled by the resident", request=request)
    return {"ok": True, "ref": doc["id"][-6:].upper()}


def _whatsapp_order_text(doc: dict, amount_paise: int) -> str:
    """The order summary a resident pastes into the committee WhatsApp group with their payment screenshot."""
    first_name = ((doc.get("name") or "").split() or [""])[0]
    out = [
        f"🍽️ ONE10 FOOD ORDER #{doc['id'][-6:].upper()}",
        f"{doc.get('tower_name') or ''} · Flat {doc.get('flat_number') or ''} · {first_name} · {_mask_mobile(doc.get('mobile'))}",
    ]
    for ln in _voucher_lines(doc):
        extra = []
        if ln["free"]:
            extra.append(f"{ln['free']} complimentary")
        if ln["kids"]:
            extra.append(f"+{ln['kids']} kid{'s' if ln['kids'] > 1 else ''} below 7, ₹0")
        out.append(f"• {ln['item']} × {ln['heads'] - ln['kids']}" + (f" ({', '.join(extra)})" if extra else ""))
    out.append(f"Amount paid by UPI: ₹{amount_paise / 100:,.0f}")
    out.append("Payment screenshot attached ✅")
    return "\n".join(out)


@router.post("/food/whatsapp-posted")
async def food_whatsapp_posted(body: dict = Body(...), request: Request = None):
    """Resident says they posted the payment screenshot + order message in the committee WhatsApp group.

    The order then never auto-expires and shows "do not pay again" until the committee records the payment.
    """
    from tokens import read_status_token
    intent_id = read_status_token((body.get("status_token") or "").strip())
    intent = await db.subscription_intents.find_one(
        {"id": intent_id}, {"_id": 0, "kind": 1, "food_subscription_id": 1}) if intent_id else None
    if not intent or intent.get("kind") != "food_subscription":
        raise HTTPException(status_code=404, detail="Order not found — open it again from My orders.")
    doc = await db.food_subscriptions.find_one_and_update(
        {"id": intent["food_subscription_id"], "payment_status": {"$in": ["pending", "partially_paid"]}},
        {"$set": {"whatsapp_posted_at": iso(), "updated_at": iso()}},
        projection=NO_ID, return_document=True,
    )
    if not doc:
        raise HTTPException(status_code=409, detail="This order is already paid, cancelled or expired.")
    await audit("food.subscription.whatsapp_posted", actor="public", entity_type="food_subscription",
                entity_id=doc["id"], request=request)
    return {"ok": True, "ref": doc["id"][-6:].upper()}


@router.post("/admin/food-subscriptions/{sid}/record-payment")
async def admin_record_food_payment(sid: str, body: dict = Body(default={}), request: Request = None,
                                    user: dict = Depends(require("receipts:manage"))):
    """Record a UPI payment the resident posted on WhatsApp (screenshot + order message).

    Never more than the amount due — overpayments are refused. A UTR can pay only one order.
    """
    from routes_collect import _insert_upi_payment, _intent_due_paise, _issue_receipt
    from vision_extract import normalize_ref
    utr = normalize_ref(body.get("utr") or "")
    if len(utr) < 6:
        raise HTTPException(status_code=400, detail="Enter the UTR / UPI reference from the screenshot.")
    try:
        amount = int(round(float(body.get("amount_rupees")) * 100))
    except (TypeError, ValueError):
        amount = 0
    if amount <= 0:
        raise HTTPException(status_code=400, detail="Enter the amount shown on the screenshot.")
    doc = await db.food_subscriptions.find_one({"id": sid}, NO_ID)
    if not doc:
        raise HTTPException(status_code=404, detail="Food order not found.")
    intent = await db.subscription_intents.find_one({"id": doc.get("intent_id")}) if doc.get("intent_id") else None
    if doc.get("payment_status") not in ("pending", "partially_paid") or not intent or \
            intent.get("status") not in ("payment_pending", "pending", "authorised", "partially_paid"):
        raise HTTPException(status_code=409, detail="This order isn't waiting for a payment.")
    due = _intent_due_paise(intent)
    if amount > due:
        raise HTTPException(status_code=409, detail=(
            f"₹{amount / 100:,.0f} is more than the ₹{due / 100:,.0f} due on this order. Overpayments aren't "
            f"accepted — check the screenshot again."))
    if await db.payments.find_one({"provider": "upi_qr", "provider_payment_id": utr}, {"_id": 1}) or \
            await db.upi_submissions.find_one({"reference_normalized": utr, "intent_id": {"$ne": intent["id"]}}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="This UTR was already used for another payment.")
    bank_seen = bool(body.get("bank_verified"))
    await _insert_upi_payment(intent_id=intent["id"], provider_payment_id=utr, amount_paise=amount,
                              proof_doc_id=None, kind="food_subscription")
    receipt = await _issue_receipt(
        intent, method="upi_qr", masked_ref="UTR-" + utr[-4:], provider_payment_id=utr, debit_account="1001",
        request=request, paid_amount_paise=amount, bank_verified=bank_seen,
        issuance_source="collection" if bank_seen else "payment_screenshot",
    )
    if not receipt:
        raise HTTPException(status_code=409, detail="Could not record — this order may have just been paid.")
    await db.receipts.update_one({"id": receipt["id"]}, {"$set": {"payment_channel": "whatsapp"}})
    if bank_seen:
        await _mark_receipts_bank_verified([receipt], user, request)
    await db.food_subscriptions.update_one({"id": sid}, {
        "$unset": {"whatsapp_posted_at": ""},
        "$set": {"whatsapp_recorded_at": iso(), "whatsapp_recorded_by": user.get("name") or user.get("login_id") or ""},
    })
    await audit("food.subscription.record_payment", actor=user, entity_type="food_subscription", entity_id=sid,
                after={"utr_last4": utr[-4:], "amount": amount, "bank_verified": bank_seen,
                       "receipt_no": receipt.get("receipt_no")}, request=request)
    return {"ok": True, "receipt_no": receipt.get("receipt_no"), "partial": bool(receipt.get("is_partial")),
            "amount_due_after_paise": int(receipt.get("amount_due_after") or 0)}


async def _food_order_card(doc: dict) -> dict:
    """One order as the buyer sees it: what it holds, what is paid, and the one thing to do next."""
    from tokens import status_token
    settlement = await _order_settlement(doc)
    status = doc.get("payment_status") or ""
    total = int(doc.get("total_amount_paise") or 0)
    paid = settlement["paid"]
    under_review = await _payment_under_review(doc)
    if status in _RELEASED_PAYMENT_STATUSES:
        state, label = status, ("Cancelled" if status == "cancelled" else "Expired — not paid")
    elif settlement["settled"]:
        state, label = "ready", ("Complimentary — voucher ready" if status == "complimentary" else "Paid — voucher ready")
    elif under_review:
        state, label = "under_review", "Payment under review — do not pay again"
    elif doc.get("whatsapp_posted_at") and status in ("pending", "partially_paid"):
        state, label = "whatsapp_posted", "Payment posted on WhatsApp — committee will verify. Do not pay again"
    elif status == "paid" or (status == "complimentary" and not settlement["settled"]):
        state, label = "problem", "Payment needs a check — contact the committee"
    elif paid > 0:
        state, label = "part_paid", f"Part-paid — ₹{(total - paid) / 100:,.0f} left to pay"
    elif status == "not_open":
        state, label = "not_open", "Payment not open yet"
    else:
        state, label = "awaiting", "Awaiting payment"
    lines = _voucher_lines(doc)
    can_pay = state in ("awaiting", "part_paid") and bool(doc.get("intent_id"))
    return {
        "id": doc["id"],
        "ref": doc["id"][-6:].upper(),
        "created_at": doc.get("created_at"),
        "name": doc.get("name") or "",
        "lines": [{"day_label": ln["day_label"], "meal_label": ln["meal_label"], "item": ln["item"],
                   "diet_label": ln["diet_label"], "heads": ln["heads"], "free": ln["free"],
                   "kids": ln["kids"]} for ln in lines],
        "heads": sum(ln["heads"] for ln in lines),
        "free_heads": sum(ln["free"] for ln in lines),
        "kid_heads": sum(ln["kids"] for ln in lines),
        "total_amount_paise": total,
        "amount_paid_paise": paid,
        "amount_due_paise": 0 if state in ("ready", "cancelled", "expired") else max(0, total - paid),
        "state": state,
        "state_label": label,
        "voucher_no": doc.get("voucher_no") if state == "ready" else None,
        "voucher_token": _food_voucher_token(doc["id"]) if state == "ready" else None,
        "status_token": status_token(doc["intent_id"]) if (can_pay or state in ("under_review", "whatsapp_posted")) else None,
        "coupons_given": bool(doc.get("coupon_id")),
        "whatsapp_text": _whatsapp_order_text(doc, max(0, total - paid)) if (can_pay or state == "whatsapp_posted") else None,
        "edit_token": None if await _resident_change_problem(doc) else food_edit_token(doc["id"]),
    }


async def _food_lookup_limit(request: Request, kind: str, flat_id: str) -> None:
    """8 lookups a minute per flat from one IP — residents on a shared Wi-Fi / mobile-carrier IP don't block each other."""
    from util import now_utc
    ip = request.client.host if request and request.client else "?"
    since = now_utc().timestamp() - 60
    if await db.receipt_lookups.count_documents({"ip": ip, "kind": kind, "flat_id": flat_id, "ts": {"$gt": since}}) >= 8:
        raise HTTPException(status_code=429, detail="Too many attempts. Please wait a minute.")
    if await db.receipt_lookups.count_documents({"ip": ip, "kind": kind, "ts": {"$gt": since}}) >= 60:
        raise HTTPException(status_code=429, detail="Too many attempts. Please wait a minute.")
    await db.receipt_lookups.insert_one({"ip": ip, "ts": now_utc().timestamp(), "kind": kind, "tower_id": "", "flat_id": flat_id or ""})


@router.post("/food/my-orders")
async def food_my_orders(body: dict = Body(...), request: Request = None):
    """Orders placed for this flat from this mobile, with pay-balance / voucher / cancel actions."""
    flat_id = (body.get("flat_id") or "").strip()
    mobile = _mobile_input(body.get("mobile"))
    if not flat_id or not mobile:
        raise HTTPException(status_code=400, detail="Select your flat and enter the 10-digit mobile used on the order.")
    await _food_lookup_limit(request, "food_my_orders", flat_id)
    await _expire_abandoned_food_orders(flat_id)
    docs = await db.food_subscriptions.find({"flat_id": flat_id}, NO_ID).sort("created_at", -1).to_list(100)
    mine = [d for d in docs if _mobile10(d.get("mobile")) == mobile]
    others_holding = [
        {"mobile_masked": _mask_mobile(d.get("mobile")), "free_heads": sum((d.get("complimentary") or {}).values())}
        for d in docs
        if _mobile10(d.get("mobile")) != mobile and d.get("complimentary")
        and d.get("payment_status") not in _RELEASED_PAYMENT_STATUSES
    ]
    return {"orders": [await _food_order_card(d) for d in mine], "others_holding": others_holding}


async def _cancel_food_order(doc: dict, *, actor, reason: str, request: Request = None) -> None:
    now = iso()
    await db.food_subscriptions.update_one({"id": doc["id"]}, {"$set": {
        "payment_status": "cancelled", "status": "cancelled", "cancelled_at": now,
        "cancel_reason": reason, "updated_at": now,
    }})
    if doc.get("intent_id"):
        await db.subscription_intents.update_one(
            {"id": doc["intent_id"], "status": {"$in": _OPEN_INTENT_STATUSES}},
            {"$set": {"status": "superseded", "superseded_at": now, "supersede_reason": "food_order_cancelled"}},
        )
    await audit("food.subscription.cancel", actor=actor, entity_type="food_subscription", entity_id=doc["id"],
                before={"payment_status": doc.get("payment_status"), "complimentary": doc.get("complimentary")},
                reason=reason, request=request)


@router.post("/admin/food-subscriptions/{sid}/cancel")
async def admin_cancel_food_order(sid: str, body: dict = Body(default={}), request: Request = None,
                                  user: dict = Depends(require("receipts:manage"))):
    """Committee cancels any order (frees its free meals; its voucher stops being valid)."""
    reason = (body.get("reason") or "").strip()
    if len(reason) < 4:
        raise HTTPException(status_code=400, detail="Give a reason for cancelling this order.")
    doc = await db.food_subscriptions.find_one({"id": sid}, NO_ID)
    if not doc:
        raise HTTPException(status_code=404, detail="Food order not found.")
    if doc.get("payment_status") in _RELEASED_PAYMENT_STATUSES:
        raise HTTPException(status_code=409, detail="This order is already cancelled or expired.")
    if doc.get("coupon_id"):
        raise HTTPException(status_code=409, detail="Coupons were already given for this order. Collect them back and Undo first.")
    async with _flat_guard(doc.get("flat_id") or sid):
        doc = await db.food_subscriptions.find_one({"id": sid}, NO_ID)
        if doc.get("payment_status") in _RELEASED_PAYMENT_STATUSES:
            raise HTTPException(status_code=409, detail="This order is already cancelled or expired.")
        if doc.get("coupon_id"):
            raise HTTPException(status_code=409, detail="Coupons were already given for this order. Collect them back and Undo first.")
        settlement = await _order_settlement(doc)
        await _cancel_food_order(doc, actor=user, reason=reason, request=request)
    return {"ok": True, "amount_paid_paise": settlement["paid"]}


async def _mark_receipts_bank_verified(receipts: list, user: dict, request: Request = None) -> None:
    now = iso()
    for r in receipts:
        await db.receipts.update_one({"id": r["id"]}, {"$set": {
            "bank_verified": True, "verification_level": "bank_statement",
            "bank_verified_at": now, "bank_verified_by": user.get("user_id") or user.get("login_id") or "",
        }})
        if r.get("payment_id"):
            await db.payments.update_many({"provider_payment_id": r["payment_id"]},
                                          {"$set": {"bank_verified": True, "bank_verified_at": now}})
        await audit("receipt.bank_verify", actor=user, entity_type="receipt", entity_id=r["id"],
                    after={"receipt_no": r.get("receipt_no"), "payment_id": r.get("payment_id")},
                    reason="Confirmed on bank statement before issuing food coupons", request=request)


async def _complimentary_problem(doc: dict) -> str:
    """Re-check an order's complimentary meals before coupons go out; '' when all is well.

    The flat must still be a subscriber, the order's free units must match its lines, and the flat's
    live orders together must stay within each quota (3 breakfasts a day, 4 Ashtami lunches, …).
    """
    from pujo_menu_2026 import COMPLIMENTARY_QUOTAS, COMPLIMENTARY_RULES, complimentary_key
    granted = {k: int(v or 0) for k, v in (doc.get("complimentary") or {}).items() if int(v or 0)}
    per_line: dict = {}
    for sel in doc.get("selections") or []:
        n = int(sel.get("complimentary_qty") or 0)
        if n:
            k = complimentary_key(sel.get("category") or sel.get("meal_code") or "",
                                  sel.get("item_day_code") or sel.get("day_code") or "")
            per_line[k] = per_line.get(k, 0) + n
    if not granted and not per_line:
        return ""
    if per_line != granted:
        return "This order's complimentary meals don't add up — check the order with the committee before giving coupons."
    if not await _flat_subscribed(doc.get("flat_id") or ""):
        return ("This flat no longer has a valid Puja subscription receipt, so its complimentary meals can't be given. "
                "Cancel this order and re-order, or restore the subscription receipt first.")
    used = await _complimentary_used(doc.get("flat_id") or "")
    labels = {r["key"]: r["label"] for r in COMPLIMENTARY_RULES}
    over = [f"{labels.get(k, k)}: {used[k]} of {q}" for k, q in COMPLIMENTARY_QUOTAS.items() if used.get(k, 0) > q]
    if over:
        return ("This flat's orders together claim more complimentary meals than allowed (" + "; ".join(over) +
                "). Cancel the extra order before giving coupons.")
    return ""


def _coupon_slips(coupon: dict) -> list:
    """One slip per person per meal, numbered {coupon_no}-001, -002, … in day / meal order."""
    slips = []
    n = 0
    for ln in coupon.get("lines") or []:
        heads = int(ln.get("heads") or 0)
        free = int(ln.get("free") or 0)
        kids = int(ln.get("kids") or 0)
        for i in range(1, heads + 1):
            n += 1
            slips.append({**ln, "serial": f"{coupon['coupon_no']}-{n:03d}", "n": i, "of": heads,
                          "is_free": i <= free, "is_kid": i > heads - kids})
    return slips


@router.post("/admin/food-subscriptions/{sid}/coupon")
async def admin_generate_coupon(sid: str, body: dict = Body(default={}), request: Request = None,
                                user: dict = Depends(require("receipts:manage"))):
    """Issue the food coupons for a settled order (once per order).

    Paid orders need every receipt matched on the bank statement first, so a fake payment screenshot
    can never turn into coupons. The issuer can confirm that check here; it is recorded against them.
    """
    doc = await db.food_subscriptions.find_one({"id": sid}, NO_ID)
    if not doc:
        raise HTTPException(status_code=404, detail="Food subscription not found.")
    if doc.get("coupon_id"):
        existing = await db.food_coupons.find_one({"id": doc["coupon_id"]}, NO_ID)
        return {"ok": True, "coupon": clean(existing), "already_existed": True}
    settlement = await _order_settlement(doc)
    if not settlement["settled"]:
        raise HTTPException(status_code=409, detail="Coupons can be given only after the full amount is paid.")
    if settlement["unverified"]:
        if not body.get("confirm_bank_verified"):
            raise HTTPException(status_code=409, detail={
                "code": "bank_unverified",
                "message": "Check these payments on the bank statement before handing out coupons.",
                "receipts": [{"receipt_no": r.get("receipt_no"), "amount_paise": r.get("total_amount"),
                              "utr": r.get("payment_id") or r.get("masked_ref") or ""}
                             for r in settlement["unverified"]],
            })
        await _mark_receipts_bank_verified(settlement["unverified"], user, request)
    problem = await _complimentary_problem(doc)
    if problem:
        raise HTTPException(status_code=409, detail=problem)
    doc = await _ensure_food_voucher(doc, settlement, actor=user, request=request)
    voucher = await _food_voucher_payload(doc, settlement)

    s = await get_settings()
    prefix = (s.get("food_subscription") or {}).get("coupon_prefix") or "ONE10-FOOD26"
    _, coupon_no = await next_formatted("food_coupon", prefix, 5)
    cid = new_id("fcoupon")
    coupon = {
        "id": cid,
        "coupon_no": coupon_no,
        "subscription_id": sid,
        "voucher_no": doc.get("voucher_no"),
        "cycle_id": doc.get("cycle_id"),
        "name": doc.get("name"),
        "mobile": doc.get("mobile"),
        "tower_name": doc.get("tower_name"),
        "flat_number": doc.get("flat_number"),
        "selections": doc.get("selections") or [],
        "lines": voucher["lines"],
        "totals": voucher["totals"],
        "slip_count": voucher["totals"]["heads"],
        "status": "issued",
        "print_count": 0,
        "issued_by": user.get("user_id"),
        "issued_by_name": user.get("name") or user.get("login_id") or "",
        "created_at": iso(),
        "updated_at": iso(),
    }
    claimed = await db.food_subscriptions.update_one(
        {"id": sid, "coupon_id": {"$in": [None, ""]}},
        {"$set": {"coupon_id": cid, "coupon_no": coupon_no, "coupons_issued_at": iso(),
                  "status": "coupon_issued", "updated_at": iso()}},
    )
    if not claimed.modified_count:
        fresh = await db.food_subscriptions.find_one({"id": sid}, NO_ID) or {}
        existing = await db.food_coupons.find_one({"id": fresh.get("coupon_id")}, NO_ID)
        return {"ok": True, "coupon": clean(existing), "already_existed": True}
    await db.food_coupons.insert_one(dict(coupon))
    await audit(
        "food.coupon.generate",
        actor=user,
        entity_type="food_coupon",
        entity_id=cid,
        after={"coupon_no": coupon_no, "subscription_id": sid, "slips": coupon["slip_count"]},
        reason="Admin generated food coupons",
        request=request,
    )
    await audit(
        "food.subscription.coupon_link",
        actor=user,
        entity_type="food_subscription",
        entity_id=sid,
        after={"coupon_no": coupon_no, "coupon_id": cid},
        request=request,
    )
    return {"ok": True, "coupon": clean(coupon), "already_existed": False}


@router.post("/admin/food-coupons/{cid}/void")
async def admin_void_food_coupon(cid: str, body: dict = Body(default={}), request: Request = None,
                                 user: dict = Depends(require("receipts:manage"))):
    """Cancel a coupon set (misprint / damaged) so a fresh set can be issued. Reason is recorded."""
    reason = (body.get("reason") or "").strip()
    if len(reason) < 4:
        raise HTTPException(status_code=400, detail="Give a reason for voiding these coupons.")
    coupon = await db.food_coupons.find_one({"id": cid}, NO_ID)
    if not coupon or coupon.get("status") == "void":
        raise HTTPException(status_code=404, detail="Active coupon set not found.")
    now = iso()
    await db.food_coupons.update_one({"id": cid}, {"$set": {
        "status": "void", "voided_at": now, "void_reason": reason, "voided_by": user.get("user_id"),
    }})
    await db.food_subscriptions.update_one(
        {"id": coupon["subscription_id"], "coupon_id": cid},
        {"$set": {"coupon_id": None, "coupon_no": None, "status": "paid", "updated_at": now},
         "$push": {"voided_coupons": {"coupon_id": cid, "coupon_no": coupon.get("coupon_no"), "voided_at": now,
                                      "reason": reason}}},
    )
    await audit("food.coupon.void", actor=user, entity_type="food_coupon", entity_id=cid,
                after={"coupon_no": coupon.get("coupon_no")}, reason=reason, request=request)
    return {"ok": True, "coupon_no": coupon.get("coupon_no")}


@router.get("/admin/food-coupons/{cid}/pdf")
async def admin_food_coupon_pdf(cid: str, request: Request = None,
                                user: dict = Depends(require("receipts:manage"))):
    """A4 sheet of per-person coupon slips — admin only. Reprints carry the same serials."""
    coupon = await db.food_coupons.find_one({"id": cid}, NO_ID)
    if not coupon:
        raise HTTPException(status_code=404, detail="Coupon not found.")
    if coupon.get("status") == "void":
        raise HTTPException(status_code=409, detail="These coupons were voided and can't be printed.")
    sub = await db.food_subscriptions.find_one({"id": coupon["subscription_id"]}, NO_ID) or {}
    if not sub or not (await _order_settlement(sub))["settled"]:
        raise HTTPException(status_code=409, detail="This order is no longer fully paid; coupons can't be printed.")
    if not coupon.get("lines"):
        coupon["lines"] = _voucher_lines(sub)
    s = await get_settings()
    pdf = food_coupon_pdf(coupon, sub, s, _coupon_slips(coupon))

    await db.food_coupons.update_one(
        {"id": cid},
        {"$inc": {"print_count": 1}, "$set": {"last_printed_at": iso(), "last_printed_by": user.get("user_id")}},
    )
    await audit(
        "food.coupon.print",
        actor=user,
        entity_type="food_coupon",
        entity_id=cid,
        after={"coupon_no": coupon.get("coupon_no"), "print_count": (coupon.get("print_count") or 0) + 1},
        reason="Admin printed food coupon slips",
        request=request,
    )
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="{coupon.get("coupon_no", "food-coupon")}.pdf"'},
    )


# =============================================================== FOOD MENU POLL / VOTING
async def _poll_tallies() -> dict:
    """Aggregate vote counts: key = day|meal|dish_id → count."""
    tallies: dict[str, int] = {}
    voter_count = 0
    diet_counts = {"pure_veg": 0, "eggetarian": 0, "non_veg": 0}
    async for vote in db.food_poll_votes.find({"is_deleted": {"$ne": True}}):
        voter_count += 1
        diet = vote.get("diet") or "pure_veg"
        if diet in diet_counts:
            diet_counts[diet] += 1
        for pick in vote.get("picks") or []:
            day = pick.get("day_code")
            meal = pick.get("meal_code")
            for dish_id in pick.get("dish_ids") or []:
                key = f"{day}|{meal}|{dish_id}"
                tallies[key] = tallies.get(key, 0) + 1
    return {"tallies": tallies, "voter_count": voter_count, "diet_counts": diet_counts}


def _proposed_menu(catalog: dict, tallies: dict, top_n: int = 6) -> list:
    """Build ultimate menu draft: top dishes per day/meal by votes (fallback to PDF suggested)."""
    proposed = []
    for day in catalog["days"]:
        meals = []
        for meal in day["meals"]:
            scored = []
            for dish in catalog["dishes"]:
                key = f"{day['code']}|{meal['code']}|{dish['id']}"
                votes = tallies.get(key, 0)
                if votes:
                    scored.append({**dish, "votes": votes})
            scored.sort(key=lambda x: (-x["votes"], x["name"]))
            if scored:
                top = scored[:top_n]
            else:
                # Fallback: pure_veg PDF suggestions
                sug = (meal.get("streams") or {}).get("pure_veg", {}).get("suggested") or []
                top = [{**d, "votes": 0} for d in sug[:top_n]]
            meals.append({
                "meal_code": meal["code"],
                "meal_label": meal["label"],
                "timing": meal.get("timing"),
                "top_dishes": top,
            })
        proposed.append({
            "day_code": day["code"],
            "day_label": day["label"],
            "date": day.get("date"),
            "weekday": day.get("weekday"),
            "meals": meals,
        })
    return proposed


@router.get("/food/poll")
async def food_poll():
    catalog = build_catalog()
    agg = await _poll_tallies()
    return {
        **catalog,
        "tallies": agg["tallies"],
        "voter_count": agg["voter_count"],
        "diet_counts": agg["diet_counts"],
        "proposed_menu": _proposed_menu(catalog, agg["tallies"]),
    }


@router.get("/food/poll/results")
async def food_poll_results():
    catalog = build_catalog()
    agg = await _poll_tallies()
    return {
        "meta": catalog["meta"],
        "voter_count": agg["voter_count"],
        "diet_counts": agg["diet_counts"],
        "tallies": agg["tallies"],
        "proposed_menu": _proposed_menu(catalog, agg["tallies"], top_n=8),
        "days": catalog["days"],
    }


@router.post("/food/poll/vote")
async def food_poll_vote(body: dict = Body(...), request: Request = None):
    """Submit or update a household vote (keyed by mobile)."""
    name = (body.get("name") or "").strip()
    mobile = "".join(c for c in str(body.get("mobile") or "") if c.isdigit())[-10:]
    diet = (body.get("diet") or "pure_veg").strip().lower()
    if diet not in {"pure_veg", "eggetarian", "non_veg"}:
        raise HTTPException(status_code=400, detail="Invalid diet preference.")
    if not name or len(mobile) != 10:
        raise HTTPException(status_code=400, detail="Name and a valid 10-digit mobile are required.")

    stream = "non_veg" if diet in {"non_veg", "eggetarian"} else "pure_veg"
    picks_in = body.get("picks") or []
    cleaned = []
    seen_slots = set()
    for raw in picks_in:
        day = (raw.get("day_code") or "").strip()
        meal = (raw.get("meal_code") or "").strip()
        if day not in _DAY_CODES or meal not in _MEAL_CODES:
            continue
        slot = f"{day}|{meal}"
        if slot in seen_slots:
            continue
        seen_slots.add(slot)
        ids = []
        for did in raw.get("dish_ids") or []:
            did = str(did).strip()
            if did in _DISH_IDS and did not in ids:
                # Eggetarian can pick veg+egg; pure_veg only veg
                dish = next(d for d in DISHES if d["id"] == did)
                if diet == "pure_veg" and dish["diet"] != "veg":
                    continue
                if diet == "eggetarian" and dish["diet"] == "nonveg":
                    continue
                ids.append(did)
            if len(ids) >= _MAX_PER_MEAL:
                break
        if ids:
            cleaned.append({"day_code": day, "meal_code": meal, "dish_ids": ids})

    if not cleaned:
        raise HTTPException(status_code=400, detail="Select at least one dish for any meal.")

    cycle_id = await get_active_cycle_id()
    now = iso()
    doc = {
        "cycle_id": cycle_id,
        "name": name,
        "mobile": mobile,
        "diet": diet,
        "stream": stream,
        "tower_id": (body.get("tower_id") or "").strip() or None,
        "tower_name": (body.get("tower_name") or "").strip() or None,
        "flat_number": (body.get("flat_number") or "").strip() or None,
        "notes": (body.get("notes") or "").strip() or None,
        "picks": cleaned,
        "updated_at": now,
        "is_deleted": False,
    }

    existing = await db.food_poll_votes.find_one({"mobile": mobile, "cycle_id": cycle_id, "is_deleted": {"$ne": True}})
    if existing:
        await db.food_poll_votes.update_one({"id": existing["id"]}, {"$set": doc})
        vote_id = existing["id"]
        created = False
    else:
        vote_id = new_id("fpoll")
        doc["id"] = vote_id
        doc["created_at"] = now
        await db.food_poll_votes.insert_one(doc)
        created = True

    await audit(
        "food.poll.vote",
        entity_type="food_poll_vote",
        entity_id=vote_id,
        after={"mobile": mobile[-4:], "diet": diet, "slots": len(cleaned), "created": created},
        request=request,
    )

    agg = await _poll_tallies()
    catalog = build_catalog()
    return {
        "ok": True,
        "vote_id": vote_id,
        "updated": not created,
        "message": "Thanks! Your votes are in — the ultimate menu updates live." if created else "Vote updated.",
        "voter_count": agg["voter_count"],
        "proposed_menu": _proposed_menu(catalog, agg["tallies"]),
    }


@router.get("/admin/food-poll/votes")
async def admin_food_poll_votes(user: dict = Depends(require("households:read", "ops:read", "receipts:read"))):
    items = await db.food_poll_votes.find({"is_deleted": {"$ne": True}}, NO_ID).sort("updated_at", -1).to_list(5000)
    agg = await _poll_tallies()
    return {"items": items, "voter_count": agg["voter_count"], "diet_counts": agg["diet_counts"]}

