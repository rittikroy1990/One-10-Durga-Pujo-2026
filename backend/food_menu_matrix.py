"""Pujo-day food menu matrix: Day × Meal × Veg/Non-veg.

Each cell maps to a food_menu_items document with matrix_key = day|meal|diet.
Excel import/export covers price, name, description, active (images via admin UI).
"""
from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from food_poll_catalog import DAYS, MEALS
from util import iso

MATRIX_DIETS = [
    {"code": "veg", "label": "Veg"},
    {"code": "non_veg", "label": "Non-veg"},
]

DEFAULT_CELL_PRICE_RUPEES = {
    "breakfast": 60,
    "lunch": 300,
    "dinner": 300,
}

TEMPLATE_HEADERS = [
    "Day Code",
    "Day Label",
    "Date",
    "Meal",
    "Diet",
    "Name",
    "Description",
    "Price (₹)",
    "Active (Y/N)",
    "Image URL",
]

HEADER_ALIASES = {
    "day_code": "day_code",
    "day": "day_code",
    "day_label": "day_label",
    "date": "menu_date",
    "menu_date": "menu_date",
    "meal": "meal",
    "category": "meal",
    "diet": "diet",
    "veg_non_veg": "diet",
    "name": "name",
    "menu_name": "name",
    "description": "description",
    "menu": "description",
    "price": "price_rupees",
    "price_": "price_rupees",
    "amount": "price_rupees",
    "active": "active",
    "active_y_n": "active",
    "image_url": "image_url",
    "image": "image_url",
}


def matrix_key(day_code: str, meal: str, diet: str) -> str:
    return f"{day_code}|{meal}|{diet}"


def parse_matrix_key(key: str) -> tuple[str, str, str] | None:
    parts = (key or "").split("|")
    if len(parts) != 3:
        return None
    return parts[0], parts[1], parts[2]


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9_]+", "_", (h or "").strip().lower()).strip("_")


def map_header(h: str) -> str | None:
    key = _norm(h)
    if key in HEADER_ALIASES:
        return HEADER_ALIASES[key]
    if key.startswith("price"):
        return "price_rupees"
    return None


def default_cell_name(day_label: str, meal_label: str, diet_label: str) -> str:
    return f"{day_label} {meal_label} ({diet_label})"


def build_skeleton() -> list[dict[str, Any]]:
    """36 cells: 6 days × 3 meals × 2 diets."""
    cells = []
    for day in DAYS:
        for meal in MEALS:
            for diet in MATRIX_DIETS:
                price = DEFAULT_CELL_PRICE_RUPEES.get(meal["code"], 300)
                cells.append({
                    "matrix_key": matrix_key(day["code"], meal["code"], diet["code"]),
                    "day_code": day["code"],
                    "day_label": day.get("short_label") or day["label"],
                    "day_full_label": day["label"],
                    "menu_date": day["date"],
                    "weekday": day.get("weekday") or "",
                    "meal": meal["code"],
                    "meal_label": meal["label"],
                    "diet": diet["code"],
                    "diet_label": diet["label"],
                    "name": default_cell_name(
                        day.get("short_label") or day["label"],
                        meal["label"],
                        diet["label"],
                    ),
                    "description": "",
                    "amount_paise": int(price * 100),
                    "amount_label": f"₹{price}",
                    "image_url": "",
                    "active": False,
                    "exists": False,
                })
    return cells


def merge_skeleton_with_db(db_items: list[dict]) -> list[dict[str, Any]]:
    by_key = {}
    for item in db_items or []:
        key = item.get("matrix_key")
        if not key and item.get("day_code") and item.get("category") and item.get("diet"):
            key = matrix_key(item["day_code"], item["category"], item["diet"])
        if key:
            by_key[key] = item

    out = []
    for cell in build_skeleton():
        existing = by_key.get(cell["matrix_key"])
        if existing:
            paise = int(existing.get("amount_paise") or cell["amount_paise"])
            out.append({
                **cell,
                "id": existing.get("id"),
                "name": existing.get("name") or cell["name"],
                "description": existing.get("description") or "",
                "amount_paise": paise,
                "amount_label": existing.get("amount_label") or f"₹{paise // 100}",
                "image_url": existing.get("image_url") or "",
                "active": bool(existing.get("active", True)),
                "exists": True,
                "updated_at": existing.get("updated_at"),
                "is_complimentary": bool(existing.get("is_complimentary")),
                "complimentary_note": existing.get("complimentary_note") or "",
                "price_note": existing.get("price_note") or "",
                "badge": existing.get("badge") or "",
            })
        else:
            out.append(cell)
    return out


def matrix_payload(db_items: list[dict]) -> dict:
    cells = merge_skeleton_with_db(db_items)
    return {
        "days": [
            {
                "code": d["code"],
                "label": d["label"],
                "short_label": d.get("short_label") or d["label"],
                "date": d["date"],
                "weekday": d.get("weekday") or "",
                "note": d.get("note") or "",
            }
            for d in DAYS
        ],
        "meals": [{"code": m["code"], "label": m["label"]} for m in MEALS],
        "diets": list(MATRIX_DIETS),
        "cells": cells,
        "columns": [
            {"meal": m["code"], "meal_label": m["label"], "diet": d["code"], "diet_label": d["label"],
             "key": f"{m['code']}|{d['code']}"}
            for m in MEALS
            for d in MATRIX_DIETS
        ],
    }


def build_matrix_template_xlsx(db_items: list[dict] | None = None) -> bytes:
    cells = merge_skeleton_with_db(db_items or [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Food menu matrix"
    header_fill = PatternFill("solid", fgColor="F5A623")
    header_font = Font(bold=True, color="3D2914")
    for col, h in enumerate(TEMPLATE_HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(cells, 2):
        ws.cell(i, 1, row["day_code"])
        ws.cell(i, 2, row["day_label"])
        ws.cell(i, 3, row["menu_date"])
        ws.cell(i, 4, row["meal"])
        ws.cell(i, 5, row["diet"])
        ws.cell(i, 6, row["name"])
        ws.cell(i, 7, row.get("description") or "")
        ws.cell(i, 8, int(row["amount_paise"]) // 100)
        ws.cell(i, 9, "Y" if row.get("active") else "N")
        ws.cell(i, 10, row.get("image_url") or "")
    widths = [12, 14, 12, 12, 10, 28, 36, 10, 12, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w
    # fix column letters properly
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    tip = wb.create_sheet("Instructions")
    tip["A1"] = "One 10 Durgotsav — Food menu matrix"
    tip["A1"].font = Font(bold=True, size=14)
    tips = [
        "Each row is one sellable cell: Day × Meal × Veg/Non-veg.",
        "Meal must be: breakfast | lunch | dinner",
        "Diet must be: veg | non_veg",
        "Price is in rupees (whole numbers preferred).",
        "Active Y = shown on public Food page; N = hidden.",
        "Image URL is optional — prefer uploading images in Admin → Food matrix.",
        "Do not change Day Code / Meal / Diet columns when re-uploading.",
        "Download template (with current data) → edit → upload to apply.",
    ]
    for i, t in enumerate(tips, 3):
        tip[f"A{i}"] = t
    tip.column_dimensions["A"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_matrix_upload(data: bytes, filename: str = "") -> list[dict]:
    """Parse xlsx/csv into normalized row dicts."""
    name = (filename or "").lower()
    rows: list[dict] = []
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        raw_rows = list(reader)
    else:
        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb[wb.sheetnames[0]]
        raw_rows = [[c if c is not None else "" for c in row] for row in ws.iter_rows(values_only=True)]

    if not raw_rows:
        return []
    headers = [map_header(str(h)) for h in raw_rows[0]]
    if not any(headers):
        raise ValueError("Could not recognise Excel headers. Use the downloaded template.")

    day_by_code = {d["code"]: d for d in DAYS}
    meal_codes = {m["code"] for m in MEALS}
    diet_codes = {d["code"] for d in MATRIX_DIETS}

    for idx, raw in enumerate(raw_rows[1:], start=2):
        if not any(str(c).strip() for c in raw if c is not None):
            continue
        mapped: dict[str, Any] = {"excel_row": idx}
        for i, field in enumerate(headers):
            if not field or i >= len(raw):
                continue
            mapped[field] = raw[i]

        day_code = str(mapped.get("day_code") or "").strip().lower()
        meal = str(mapped.get("meal") or "").strip().lower().replace("-", "_")
        diet = str(mapped.get("diet") or "").strip().lower().replace("-", "_").replace(" ", "_")
        if diet in ("nonveg", "non-veg", "nv"):
            diet = "non_veg"
        if diet in ("vegetarian", "veggie"):
            diet = "veg"

        if day_code not in day_by_code:
            mapped["error"] = f"Unknown day code: {day_code}"
            rows.append(mapped)
            continue
        if meal not in meal_codes:
            mapped["error"] = f"Meal must be breakfast/lunch/dinner (got {meal})"
            rows.append(mapped)
            continue
        if diet not in diet_codes:
            mapped["error"] = f"Diet must be veg or non_veg (got {diet})"
            rows.append(mapped)
            continue

        day = day_by_code[day_code]
        price_raw = mapped.get("price_rupees")
        try:
            price = float(str(price_raw).replace(",", "").strip() or 0)
        except ValueError:
            mapped["error"] = "Invalid price"
            rows.append(mapped)
            continue
        if price < 0:
            mapped["error"] = "Price cannot be negative"
            rows.append(mapped)
            continue

        active_raw = str(mapped.get("active") if mapped.get("active") is not None else "Y").strip().lower()
        active = active_raw in ("1", "y", "yes", "true", "on", "active")

        name_val = str(mapped.get("name") or "").strip()
        meal_label = next(m["label"] for m in MEALS if m["code"] == meal)
        diet_label = next(d["label"] for d in MATRIX_DIETS if d["code"] == diet)
        if not name_val:
            name_val = default_cell_name(day.get("short_label") or day["label"], meal_label, diet_label)

        rows.append({
            "excel_row": idx,
            "matrix_key": matrix_key(day_code, meal, diet),
            "day_code": day_code,
            "day_label": day.get("short_label") or day["label"],
            "menu_date": day["date"],
            "meal": meal,
            "diet": diet,
            "name": name_val[:120],
            "description": str(mapped.get("description") or "").strip()[:800],
            "price_rupees": price,
            "amount_paise": int(round(price * 100)),
            "active": active,
            "image_url": str(mapped.get("image_url") or "").strip()[:500],
        })
    return rows


_DAY_ALIASES = {
    "sasthi": "sasthi",
    "maha sasthi": "sasthi",
    "shashthi": "sasthi",
    "shasthi": "sasthi",
    "saptami": "saptami",
    "maha saptami": "saptami",
    "saptami ashtami": "saptami_ashtami",
    "saptami / ashtami": "saptami_ashtami",
    "sap-asht": "saptami_ashtami",
    "sap–asht": "saptami_ashtami",
    "ashtami": "ashtami",
    "maha ashtami": "ashtami",
    "nabami": "nabami",
    "navami": "nabami",
    "maha nabami": "nabami",
    "maha navami": "nabami",
    "dashami": "dashami",
    "vijaya dashami": "dashami",
    "bijoya dashami": "dashami",
}

_MEAL_ALIASES = {
    "breakfast": "breakfast",
    "bf": "breakfast",
    "morning": "breakfast",
    "lunch": "lunch",
    "afternoon": "lunch",
    "bhog": "lunch",
    "dinner": "dinner",
    "night": "dinner",
    "evening": "dinner",
}


def _detect_day(line: str) -> str | None:
    low = re.sub(r"\s+", " ", (line or "").strip().lower())
    for alias, code in sorted(_DAY_ALIASES.items(), key=lambda x: -len(x[0])):
        if alias in low:
            return code
    return None


def _detect_meal(line: str) -> str | None:
    low = re.sub(r"\s+", " ", (line or "").strip().lower())
    for alias, code in _MEAL_ALIASES.items():
        if re.search(rf"\b{re.escape(alias)}\b", low):
            return code
    return None


def _detect_diet(line: str) -> str | None:
    low = re.sub(r"\s+", " ", (line or "").strip().lower())
    for alias in ("non-vegetarian", "non vegetarian", "non-veg", "non veg", "nonveg"):
        if alias in low:
            return "non_veg"
    for alias in ("pure vegetarian", "pure veg", "vegetarian", "veg"):
        if re.search(rf"\b{re.escape(alias)}\b", low):
            return "veg"
    return None


def _docx_paragraphs(data: bytes) -> list[str]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            joined = " | ".join(c for c in cells if c)
            if joined:
                lines.append(joined)
    return lines


def parse_overall_menu_text(lines: list[str]) -> list[dict]:
    """Heuristic Day → Meal → Diet → dish lines parser for Word/plain menus."""
    cur_day = None
    cur_meal = None
    cur_diet = None
    buckets: dict[str, list[str]] = {}

    header_only = re.compile(
        r"(?i)^[\s\-–—:|]*(breakfast|lunch|dinner|veg|non[-\s]?veg|"
        r"vegetarian|non[-\s]?vegetarian|pure\s+veg)[\s\-–—:|]*$"
    )

    for raw in lines:
        line = (raw or "").strip()
        if not line or len(line) < 2:
            continue

        day = _detect_day(line)
        meal = _detect_meal(line)
        diet = _detect_diet(line)

        # Day heading (optionally with date); reset meal/diet
        if day and meal is None and len(line) < 80:
            cur_day = day
            cur_meal = None
            cur_diet = None
            continue

        if day and meal is None:
            cur_day = day

        if meal and len(line) < 60:
            cur_meal = meal
            if diet:
                cur_diet = diet
            if ":" in line:
                after = line.split(":", 1)[1].strip()
                if after and cur_day and cur_meal and cur_diet:
                    buckets.setdefault(matrix_key(cur_day, cur_meal, cur_diet), []).append(after)
                continue
            if header_only.match(line) or diet is not None:
                continue

        if diet and meal is None and len(line) < 48:
            cur_diet = diet
            if ":" in line:
                after = line.split(":", 1)[1].strip()
                if after and cur_day and cur_meal and cur_diet:
                    buckets.setdefault(matrix_key(cur_day, cur_meal, cur_diet), []).append(after)
            continue

        if cur_day and cur_meal and cur_diet:
            if re.fullmatch(r"₹?\s*\d+([.,]\d+)?\s*(/-)?", line):
                continue
            if header_only.match(line):
                continue
            buckets.setdefault(matrix_key(cur_day, cur_meal, cur_diet), []).append(line)

    rows: list[dict] = []
    for day in DAYS:
        for meal in MEALS:
            for diet in MATRIX_DIETS:
                key = matrix_key(day["code"], meal["code"], diet["code"])
                dishes = buckets.get(key) or []
                desc = "; ".join(dishes)[:800] if dishes else ""
                price = DEFAULT_CELL_PRICE_RUPEES.get(meal["code"], 300)
                name = default_cell_name(
                    day.get("short_label") or day["label"],
                    meal["label"],
                    diet["label"],
                )
                rows.append({
                    "matrix_key": key,
                    "day_code": day["code"],
                    "day_label": day.get("short_label") or day["label"],
                    "menu_date": day["date"],
                    "meal": meal["code"],
                    "diet": diet["code"],
                    "name": name,
                    "description": desc,
                    "price_rupees": price,
                    "amount_paise": int(price * 100),
                    "active": bool(dishes),
                    "image_url": "",
                    "filled": bool(dishes),
                })
    return rows


def parse_overall_menu_file(data: bytes, filename: str = "") -> dict:
    """Parse overall menu upload. Returns {kind, rows, note, text_preview?}."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xls", ".csv")):
        rows = parse_matrix_upload(data, filename)
        return {"kind": "matrix_excel", "rows": rows, "note": "Excel matrix import"}
    if name.endswith(".docx"):
        lines = _docx_paragraphs(data)
        rows = parse_overall_menu_text(lines)
        filled = sum(1 for r in rows if r.get("filled"))
        return {
            "kind": "docx",
            "rows": rows,
            "text_preview": "\n".join(lines[:40]),
            "note": f"Word menu parsed into {filled} filled cells",
        }
    if name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        rows = parse_overall_menu_text(lines)
        filled = sum(1 for r in rows if r.get("filled"))
        return {"kind": "text", "rows": rows, "note": f"Text menu parsed into {filled} filled cells"}
    if name.endswith((".pdf", ".png", ".jpg", ".jpeg", ".webp", ".gif")):
        return {
            "kind": "document",
            "rows": [],
            "note": "Stored as overall menu document. Use .docx or Excel to fill day × meal cells automatically.",
        }
    raise ValueError("Unsupported file. Upload .docx, .xlsx, .csv, .pdf or an image.")
