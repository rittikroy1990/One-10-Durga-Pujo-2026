"""Accounting (double-entry ledger, trial balance) and bank/gateway reconciliation."""
import csv
import hashlib
import io

from fastapi import APIRouter, Depends, Request, HTTPException, Body, UploadFile, File, Form

from db import db, new_id, clean
from config import get_settings
from util import iso, now_utc, fmt_inr
from audit import audit
from ledger import post_journal, reverse_journal, trial_balance, account_balance, LedgerError
from auth import require

router = APIRouter(prefix="/api")


# --------------------------------------------------------------- Accounting
@router.get("/accounting/chart")
async def chart(user: dict = Depends(require("accounting:read"))):
    accs = await db.chart_of_accounts.find({}, {"_id": 0}).sort("code", 1).to_list(200)
    return {"items": accs}


@router.get("/accounting/journals")
async def journals(user: dict = Depends(require("accounting:read"))):
    js = await db.journal_entries.find({}, {"_id": 0}).sort("posted_at", -1).to_list(5000)
    return {"items": js, "count": len(js)}


@router.get("/accounting/journals/{jid}")
async def journal_detail(jid: str, user: dict = Depends(require("accounting:read"))):
    j = await db.journal_entries.find_one({"id": jid}, {"_id": 0})
    if not j:
        raise HTTPException(status_code=404, detail="Not found.")
    return j


@router.post("/accounting/journals")
async def post_manual_journal(body: dict = Body(...), request: Request = None,
                              user: dict = Depends(require("accounting:post"))):
    lines = body.get("lines", [])
    try:
        j = await post_journal(source_type="manual", source_id=body.get("source_id", ""),
                               narration=body.get("narration", "Manual journal"),
                               lines=lines, actor=user, approved_by=body.get("approved_by", ""))
    except LedgerError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return j


@router.post("/accounting/journals/{jid}/reverse")
async def reverse(jid: str, body: dict = Body(...), request: Request = None,
                  user: dict = Depends(require("accounting:post"))):
    try:
        rev = await reverse_journal(jid, reason=body.get("reason", "correction"), actor=user)
    except LedgerError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return rev


@router.get("/accounting/trial-balance")
async def tb(user: dict = Depends(require("accounting:read"))):
    rows = await trial_balance()
    return {"items": rows, "total_debit": sum(r["debit"] for r in rows),
            "total_credit": sum(r["credit"] for r in rows)}


# --------------------------------------------------------------- Reconciliation
@router.get("/recon/bank-accounts")
async def bank_accounts(user: dict = Depends(require("recon:read"))):
    accs = await db.bank_accounts.find({}, {"_id": 0}).to_list(50)
    return {"items": accs}


def _parse_rows(content: bytes, filename: str):
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = content.decode("utf-8", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    return rows


@router.post("/recon/import")
async def recon_import(file: UploadFile = File(...), amount_col: str = Form(...),
                       date_col: str = Form(...), ref_col: str = Form(...),
                       type_col: str = Form(""), bank_account_id: str = Form("bank_main"),
                       request: Request = None, user: dict = Depends(require("recon:manage"))):
    content = await file.read()
    content_hash = hashlib.sha256(content).hexdigest()
    rows = _parse_rows(content, file.filename)
    if not rows:
        raise HTTPException(status_code=400, detail="Empty file.")
    header = [str(h).strip() for h in rows[0]]

    def idx(col):
        try:
            return header.index(col)
        except ValueError:
            try:
                return int(col)
            except Exception:
                raise HTTPException(status_code=400, detail=f"Column '{col}' not found in header.")

    ai, di, ri = idx(amount_col), idx(date_col), idx(ref_col)
    ti = idx(type_col) if type_col else None

    imp_id = new_id("imp")
    total = 0
    line_docs = []
    for r in rows[1:]:
        if not any(str(x).strip() for x in r):
            continue
        try:
            amt = int(round(float(str(r[ai]).replace(",", "").strip() or 0) * 100))
        except Exception:
            amt = 0
        line = {"id": new_id("bl"), "import_id": imp_id, "bank_account_id": bank_account_id,
                "amount_paise": amt, "date": str(r[di]).strip(), "ref": str(r[ri]).strip(),
                "type": (str(r[ti]).strip() if ti is not None and ti < len(r) else "credit"),
                "matched": False, "match_id": "", "created_at": iso()}
        line_docs.append(line)
        total += amt
    if line_docs:
        await db.bank_statement_lines.insert_many([dict(x) for x in line_docs])
    # store original file + hash
    try:
        from storage import save_document
        await save_document(data=content, filename=file.filename, doc_type="bank_statement_imports",
                            linked_type="bank_statement_import", linked_id=imp_id,
                            uploaded_by=user["user_id"])
    except Exception:
        pass
    imp = {"id": imp_id, "bank_account_id": bank_account_id, "filename": file.filename,
           "content_hash": content_hash, "line_count": len(line_docs), "total_paise": total,
           "mapping": {"amount": amount_col, "date": date_col, "ref": ref_col, "type": type_col},
           "imported_by": user["user_id"], "created_at": iso()}
    await db.bank_statement_imports.insert_one(dict(imp))
    await audit("recon.import", actor=user, entity_type="bank_statement_import", entity_id=imp_id,
                after={"lines": len(line_docs), "total": total, "hash": content_hash}, request=request)
    return {"import": clean(imp), "lines": len(line_docs)}


@router.get("/recon/lines")
async def recon_lines(user: dict = Depends(require("recon:read"))):
    lines = await db.bank_statement_lines.find({}, {"_id": 0}).sort("created_at", -1).to_list(5000)
    return {"items": lines}


@router.get("/recon/suggestions")
async def recon_suggestions(user: dict = Depends(require("recon:read"))):
    """Suggest matches: exact (amount + ref token) vs suggested (amount only)."""
    settings = await get_settings()
    threshold = settings["approval_thresholds_paise"].get("recon_auto_match_confidence", 0.9)
    lines = await db.bank_statement_lines.find({"matched": False}, {"_id": 0}).to_list(2000)
    receipts = await db.receipts.find({"status": "issued"}, {"_id": 0}).to_list(5000)
    suggestions = []
    for ln in lines:
        for r in receipts:
            if r["total_amount"] == ln["amount_paise"]:
                ref = (ln.get("ref") or "").lower()
                token = (r.get("masked_ref") or "").lower().replace("xxxx", "").replace("utr-", "")
                confidence = 0.95 if token and token[-4:] in ref else 0.6
                suggestions.append({"line_id": ln["id"], "receipt_id": r["id"],
                                    "receipt_no": r["receipt_no"], "amount": ln["amount_paise"],
                                    "bank_ref": ln.get("ref"), "confidence": confidence,
                                    "state": "exact" if confidence >= threshold else "suggested"})
    return {"items": suggestions, "threshold": threshold}


@router.post("/recon/match")
async def recon_match(body: dict = Body(...), request: Request = None,
                      user: dict = Depends(require("recon:manage"))):
    line = await db.bank_statement_lines.find_one({"id": body.get("line_id")})
    receipt = await db.receipts.find_one({"id": body.get("receipt_id")})
    if not line or not receipt:
        raise HTTPException(status_code=404, detail="Line or receipt not found.")
    confidence = float(body.get("confidence") or 1.0)
    settings = await get_settings()
    threshold = settings["approval_thresholds_paise"].get("recon_auto_match_confidence", 0.9)
    state = "exact" if confidence >= threshold else "suggested"
    if state == "suggested" and not body.get("human_approved"):
        raise HTTPException(status_code=400, detail="Suggested match below threshold requires human approval.")
    match = {"id": new_id("match"), "line_id": line["id"], "receipt_id": receipt["id"],
             "amount_paise": line["amount_paise"], "confidence": confidence, "state": state,
             "approved_by": user["user_id"], "created_at": iso()}
    await db.reconciliation_matches.insert_one(dict(match))
    await db.bank_statement_lines.update_one({"id": line["id"]}, {"$set": {"matched": True, "match_id": match["id"]}})
    await audit("recon.match", actor=user, entity_type="reconciliation_match", entity_id=match["id"],
                after={"receipt": receipt["receipt_no"], "state": state}, request=request)
    return clean(match)


@router.get("/recon/dashboard")
async def recon_dashboard(user: dict = Depends(require("recon:read"))):
    stale_before = now_utc().timestamp() - 3 * 3600
    orders = await db.payment_orders.find({"status": "created"}, {"_id": 0}).to_list(2000)
    stale = [o for o in orders]
    unident = await db.bank_statement_lines.find({"matched": False, "type": {"$ne": "debit"}},
                                                 {"_id": 0}).to_list(2000)
    return {
        "duplicate_payments": await db.duplicate_payments.find({}, {"_id": 0}).to_list(500),
        "amount_mismatch": await db.payment_orders.find({"status": "reconciliation_required"}, {"_id": 0}).to_list(500),
        "stale_pending_orders": stale,
        "refunds_processing": await db.refunds.find({"status": "processing"}, {"_id": 0}).to_list(500),
        "cash_not_deposited": await db.cash_collections.find({"status": "accepted", "deposited": {"$ne": True}}, {"_id": 0}).to_list(500),
        "cheques_pending": await db.cheques.find({"status": "pending_clearing"}, {"_id": 0}).to_list(500),
        "bank_credits_unidentified": unident,
    }


@router.get("/recon/balances")
async def recon_balances(user: dict = Depends(require("recon:read"))):
    out = []
    names = {"1001": "Bank", "1002": "Cash on hand", "1003": "Gateway clearing"}
    bank = await db.bank_accounts.find_one({"id": "bank_main"})
    opening = bank.get("opening_balance_paise", 0) if bank else 0
    for code, label in names.items():
        bal = await account_balance(code)
        out.append({"code": code, "account": label,
                    "opening": opening if code == "1001" else 0,
                    "movement": bal, "closing": (opening if code == "1001" else 0) + bal,
                    "closing_display": fmt_inr((opening if code == "1001" else 0) + bal)})
    return {"items": out}
